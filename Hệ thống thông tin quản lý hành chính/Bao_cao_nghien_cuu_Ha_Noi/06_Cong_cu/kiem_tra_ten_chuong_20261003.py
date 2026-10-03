from pathlib import Path
from zipfile import ZipFile
from hashlib import sha256
from docx import Document
from PIL import Image, ImageDraw, ImageChops
import json, sys
import ket_xuat_va_doi_chieu_oo as qa

ROOT = Path(__file__).resolve().parents[1]
Q = ROOT / '05_Doi_chieu/Ten_chuong_20261003'
DOC = ROOT / '04_Bao_cao/Phan_tich_thiet_ke_huong_doi_tuong_Ha_Noi.docx'
qa.Q = Q
if sys.argv[1] == 'restore':
    qa.restore(DOC)
    print('Restored original style parts')
    raise SystemExit

pages = qa.render(Q / 'Bao_cao.pdf', Q / 'Trang')
qa.check(DOC)
old = Document(Q / 'Ban_truoc_doi_ten.docx')
new = Document(DOC)
def body(d):
    out = []
    start = False
    for p in d.paragraphs:
        if p.style.name == 'Heading 1' and p.text.startswith('CHƯƠNG 1.'):
            start = True
        if start and not (p.style.name == 'Heading 1' and p.text.startswith('CHƯƠNG ')):
            out.append(p.text)
    return out
def tables(d):
    return [tuple(tuple(c.text for c in row.cells) for row in t.rows) for t in d.tables]
plan = json.loads((Q / 'Doi_ten_chuong.json').read_text('utf-8'))
titles = [p.text for p in new.paragraphs if p.style.name == 'Heading 1' and p.text.startswith('CHƯƠNG ')]
with ZipFile(DOC) as z, ZipFile(Q / 'Ban_truoc_doi_ten.docx') as before:
    protected = [n for n in before.namelist() if n.startswith(('word/media/', 'word/styles', 'word/header', 'word/footer', 'word/numbering', 'word/theme/', 'word/fontTable'))]
    same_parts = all(z.read(n) == before.read(n) for n in protected)
checks = {
    'Tên bốn chương đúng phương án': titles == [f'CHƯƠNG {i}. {t}' for i,t in plan['titles'].items()],
    'Nội dung từ chương đầu đến phụ lục giữ nguyên': body(old) == body(new),
    'Giữ nguyên nội dung 98 bảng': tables(old) == tables(new) and len(new.tables) == 98,
    'Hình ảnh và định dạng giữ nguyên': same_parts,
    'Không có chữ vượt trang hoặc lỗi trường': not any(p['overflow'] or 'Error!' in p['text'] or 'Lỗi!' in p['text'] for p in pages),
}
folder = Q / 'Trang_doi_chieu'; folder.mkdir(exist_ok=True)
images = sorted((Q / 'Trang').glob('trang_*.png'))
for i in range(0,len(images),4):
    ims = [Image.open(x).convert('RGB') for x in images[i:i+4]]
    w,h = ims[0].size
    sheet = Image.new('RGB',(w*2,(h+30)*2),'#dddddd'); draw = ImageDraw.Draw(sheet)
    for j,im in enumerate(ims):
        x=j%2*w; y=j//2*(h+30)
        draw.text((x+15,y+5),f'TRANG {i+j+1}',fill='black'); sheet.paste(im,(x,y+30))
    sheet.save(folder / f'bo_{i//4+1:02}.png')
changed = []
for i,img in enumerate(images):
    prev = ROOT / '05_Doi_chieu/Gop_chuong_20261003/Trang' / img.name
    if not prev.exists() or ImageChops.difference(Image.open(img).convert('RGB'),Image.open(prev).convert('RGB')).getbbox():
        changed.append(i+1)
result = {'docx':str(DOC),'sha256':sha256(DOC.read_bytes()).hexdigest(),'pages':len(pages),'titles':titles,'checks':checks,'changed_visual_pages':changed}
(Q / 'Ket_qua_kiem_tra.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),'utf-8')
print(json.dumps(result,ensure_ascii=False,indent=2))
if not all(checks.values()): raise SystemExit(1)
