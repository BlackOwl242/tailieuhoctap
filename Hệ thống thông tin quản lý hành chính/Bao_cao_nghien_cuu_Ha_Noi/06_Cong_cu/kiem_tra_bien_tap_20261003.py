from pathlib import Path
import sys,json,re,hashlib
from PIL import Image,ImageDraw
from docx import Document
from lxml import etree as E
import ket_xuat_va_doi_chieu_oo as qa
ROOT=Path(__file__).resolve().parents[1];Q=ROOT/'05_Doi_chieu/Bien_tap_lai_20261003';qa.Q=Q
DOC=ROOT/'04_Bao_cao/Phan_tich_thiet_ke_huong_doi_tuong_Ha_Noi.docx'
if sys.argv[1]=='restore':qa.restore(DOC)
elif sys.argv[1]=='render':
 qa.render(Q/'Bao_cao.pdf',Q/'Trang');pages=sorted((Q/'Trang').glob('trang_*.png'));folder=Q/'Trang_doi_chieu';folder.mkdir(exist_ok=True)
 for i in range(0,len(pages),4):
  ims=[Image.open(p) for p in pages[i:i+4]];w,h=ims[0].size;c=Image.new('RGB',(w*2,(h+30)*2),'#dddddd');d=ImageDraw.Draw(c)
  for j,im in enumerate(ims):
   x=j%2*w;y=j//2*(h+30);d.text((x+15,y+5),'TRANG '+str(i+j+1),fill='black');c.paste(im,(x,y+30))
  c.save(folder/f'bo_{i//4+1:02}.png')
 print('Đã tạo',len(pages),'trang và',len(list(folder.glob('bo*.png'))),'bộ ảnh.')
elif sys.argv[1]=='check':
 qa.check(DOC);d=Document(DOC);text='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells);text=text.replace('\u200b','').replace('\u00ad','')
 specs=[t for t in d.tables if len(t.rows)>1 and t.cell(1,0).text=='Mã và mục tiêu']
 checks={'36 đặc tả có mục tiêu riêng':len(specs)==36 and all(re.match(r'UC\d{2}\. ',t.cell(1,1).text) for t in specs),'Luồng chính trình bày từng bước':all(len(t.cell(5,1).paragraphs)>=4 for t in specs),'Không còn mã UCN hoặc R':not re.search(r'\b(?:UCN|R)\d+\b',text),'Không nhắc mẫu trong báo cáo':not any(x in text.lower() for x in ['pttk_oop_hr','talentconnect','tài liệu mẫu','audit mẫu']),'Đủ 10 nguồn':set(re.findall(r'\[(\d+)\]',text))==set(map(str,range(1,11))),'Không còn nhãn khó hiểu của sơ đồ cũ':all(x not in text for x in ['Đủ điều kiện đã đặc tả','Kết quả đã xác nhận và nhật ký được lưu cùng giao dịch; thử lại không ghi trùng'])}
 figures=[]
 for n in range(1,37):
  s=(ROOT/f'03_Thiet_ke/Huong_doi_tuong/So_do/UC{n:02}_trinh_tu.svg').read_text('utf-8');visible=' '.join(E.fromstring(s.encode()).xpath('//*[local-name()="text"]/text()'))
  figures.append({'uc':f'UC{n:02}','has_alternative':'alt' in visible,'has_specific_labels':'Kết quả và phiên bản' not in visible and 'dt :' not in visible,'uses_arial':'font-family="Arial"' in s})
 checks['36 biểu đồ trình tự có nhánh và nhãn nghiệp vụ']=all(all(v for k,v in f.items() if k!='uc') for f in figures)
 pages=json.loads((Q/'Trang/van_ban_theo_trang.json').read_text('utf-8'));checks['Không có chữ vượt trang']=all(not p['overflow'] for p in pages)
 result={'sha256':hashlib.sha256(DOC.read_bytes()).hexdigest(),'pages':len(pages),'checks':checks,'figures':figures,'scope':'Kiểm tra tài liệu và thiết kế; chưa kiểm thử phần mềm thực tế.'}
 (Q/'Kiem_tra_ban_cuoi.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(result['checks'],ensure_ascii=False,indent=2))
