"""Insert current UI documentation without regenerating the research chapters."""
from copy import deepcopy
from pathlib import Path
from PIL import Image
from docx.shared import Inches, Pt, Twips, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from noi_dung_bao_cao import ROOT, DATE, ROLE_ROWS, FLOW_ROWS, OVERVIEW, UI_GROUP_ROWS, CMS_ROWS, COVERAGE, UI_TITLES, IMPLEMENTATION_NOTES, manifest


class Writer:
    def __init__(self, doc, anchor=None):
        self.doc, self.anchor = doc, anchor

    def place(self, el):
        if self.anchor is not None:
            self.anchor.addprevious(el)

    def p(self, text='', style=None):
        p = self.doc.add_paragraph(text, style)
        self.place(p._p)
        if style is None:
            p.paragraph_format.keep_with_next = False
        return p

    def caption(self, kind, chapter, number, title):
        p = self.p()
        p.paragraph_format.first_line_indent = Twips(0)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if kind == 'Hình' else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.keep_with_next = kind == 'Bảng'
        p.add_run(f'{kind} {chapter}.')
        for tag, text in [('begin',None),('instruction',' SEQ '+kind+' '),('separate',None),('result',str(number)),('end',None)]:
            r = p.add_run()
            if tag == 'instruction':
                el = OxmlElement('w:instrText'); el.set(qn('xml:space'),'preserve'); el.text = text; r._r.append(el)
            elif tag == 'result':
                r.add_text(text)
            else:
                el=OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'),tag); r._r.append(el)
        p.add_run(': '+title)
        for r in p.runs:
            r.font.name='Times New Roman'; r.font.size=Pt(11); r.font.color.rgb=RGBColor(0,0,0)
            r.italic=kind=='Hình'; r.bold=kind=='Bảng'
        return p

    def table(self, rows, widths):
        t = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.autofit=False
        self.place(t._tbl)
        pr=t._tbl.tblPr
        borders=OxmlElement('w:tblBorders')
        for edge in ('top','left','bottom','right','insideH','insideV'):
            e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
        pr.append(borders)
        for col,w in zip(t.columns,widths): col.width=Twips(w)
        for i,row in enumerate(t.rows):
            trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
            if i==0: trpr.append(OxmlElement('w:tblHeader'))
            for j,cell in enumerate(row.cells):
                cell.width=Twips(widths[j]);cell.text=str(rows[i][j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr=cell._tc.get_or_add_tcPr()
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'203A63' if i==0 else 'F3F6FA' if i%2 else 'FFFFFF');tcpr.append(sh)
                margins=OxmlElement('w:tcMar')
                for edge in ('top','left','bottom','right'):
                    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
                tcpr.append(margins)
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent=Twips(0);p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.05;p.paragraph_format.keep_with_next=False
                    p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 or (j==0 and len(cell.text)<10) else WD_ALIGN_PARAGRAPH.LEFT
                    for r in p.runs:
                        r.font.name='Times New Roman';r.font.size=Pt(11);r.bold=i==0;r.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
        return t

    def section_break(self, properties):
        p=self.p();p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(1);p.add_run().font.size=Pt(1)
        sect=deepcopy(properties)
        kind=sect.find(qn('w:type'))
        if kind is None:kind=OxmlElement('w:type');sect.insert(0,kind)
        kind.set(qn('w:val'),'nextPage')
        p._p.get_or_add_pPr().append(sect)


def add_ui_section(doc, anchor=None):
    w=Writer(doc,anchor)
    captures=manifest()['captures']
    for text in OVERVIEW:w.p(text)
    w.caption('Bảng','4',8,'Trách nhiệm và kết quả công việc của từng vai trò')
    w.table([['Vai trò','Công việc','Kết quả bàn giao'],*ROLE_ROWS],[1800,3650,3622])
    w.caption('Bảng','4',9,'Luồng hồ sơ và dữ liệu giữa các vai trò')
    w.table([['Bước','Người thực hiện','Công việc','Thông tin được lưu'],*FLOW_ROWS],[650,1500,2300,4622])
    w.caption('Bảng','4',10,'Các nhóm giao diện của hệ thống')
    w.table([['Khu vực','Nhóm chức năng','Màn hình chính'],*UI_GROUP_ROWS],[1700,2100,5272])
    w.p('Cẩm nang sử dụng hướng dẫn từng công việc theo vai trò và có liên kết đến màn hình tương ứng. Cán bộ đăng nhập trước khi thực hiện nghiệp vụ. Trang tài khoản cho phép xem quyền được cấp, cập nhật thông tin liên hệ, đổi mật khẩu và kết thúc các phiên đăng nhập khác.')
    w.p('Đối với tin tức, quản trị viên biên tập bài, chọn chuyên mục, hình ảnh và nguồn tham khảo trước khi xuất bản. Hệ thống lưu phiên bản bài viết để theo dõi thay đổi và xử lý trường hợp nhiều thao tác cập nhật cùng lúc.')
    w.caption('Bảng','4',11,'Trạng thái và kiểm soát tin tức')
    w.table([['Trạng thái hoặc thao tác','Người thực hiện','Kết quả'],*CMS_ROWS],[2250,1800,5022])
    w.p(f'Các hình dưới đây minh họa giao diện ứng dụng do em xây dựng, được chụp ngày {DATE}. Số liệu xuất hiện trong ảnh thuộc cơ sở dữ liệu của ứng dụng.')
    order=[1,2,17,18,19,42,43,44,45,20,21,22,12,3,4,5,30,6,37,38,39,40,8,10,11,15,16,31,7,32,33,34,9,35,13,14,36,41,23,24,25,26,27,28,29]
    by_num={int(r['file'].split('_')[0]):r for r in captures}
    captions=[]
    for i,n in enumerate(order):
        record=by_num[n];file=ROOT/'03_Thiet_ke/Anh_giao_dien_he_thong'/record['file']
        with Image.open(file) as im:px,py=im.size
        factor=min(6.25/px,7.3/py)
        p=w.p();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Twips(0);p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1;p.paragraph_format.keep_with_next=True
        pic=p.add_run().add_picture(str(file),width=Inches(px*factor),height=Inches(py*factor))
        title=UI_TITLES[n]
        pic._inline.docPr.set('descr',title+'; ảnh chụp ngày '+DATE)
        w.caption('Hình','4',26+i,title)
        captions.append(f'Hình 4.{26+i}: {title}')
    return captions


def add_coverage_section(doc,anchor=None):
    w=Writer(doc,anchor)
    w.p('6.5. Đánh giá mức độ đáp ứng các ca sử dụng', 'Heading 2')
    w.p('Kết quả triển khai được đối chiếu với 36 ca sử dụng đã phân tích ở chương 3. Bảng 6.40 nêu chức năng đã thực hiện được và những nội dung cần tiếp tục bổ sung. Các ca sử dụng liên quan đến dịch vụ bên ngoài được đánh giá theo tình trạng kết nối của ứng dụng.')
    w.caption('Bảng','6',40,'Kết quả triển khai theo ca sử dụng')
    w.table([['Ca sử dụng','Kết quả','Chức năng','Nội dung đáp ứng và phần cần bổ sung'],*[[r['id']+' '+r['name'],r['status'],r['ui'],r['note']] for r in COVERAGE]],[1800,1450,1850,3972])


def add_implementation_section(doc,anchor=None):
    w=Writer(doc,anchor)
    w.p('6.4. Kết quả triển khai và hạn chế','Heading 2')
    for text in IMPLEMENTATION_NOTES:
        w.p(text)
