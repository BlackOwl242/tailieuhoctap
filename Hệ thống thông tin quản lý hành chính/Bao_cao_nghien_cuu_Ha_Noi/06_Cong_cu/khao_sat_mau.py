from pathlib import Path
from collections import Counter
from zipfile import ZipFile
from hashlib import sha256
import json, sys
from docx import Document
from lxml import etree

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parents[1]
REF = ROOT.parents[1] / 'Công nghệ phần mềm' / 'Bao_Cao_Bai_Tap_Lon_CNPM_TalentConnect.docx'
QA = ROOT / '05_Doi_chieu'
QA.mkdir(parents=True, exist_ok=True)
d=Document(REF)
def inner(e):
    if e is None: return None
    return ''.join(etree.tostring(x, encoding='unicode') for x in e)
info={'reference':str(REF),'sha256':sha256(REF.read_bytes()).hexdigest(),
      'sections':[inner(s._sectPr) for s in d.sections],
      'styles':{s.name:inner(s.element) for s in d.styles},
      'paragraphs':[{'index':i,'style':p.style.name,'text':p.text,'properties':inner(p._p.pPr),
                     'runs':[{'text':r.text,'properties':inner(r._r.rPr)} for r in p.runs]} for i,p in enumerate(d.paragraphs)],
      'tables':[{'properties':inner(t._tbl.tblPr),'rows':len(t.rows),'columns':len(t.columns),
                 'head':[c.text for c in t.rows[0].cells]} for t in d.tables]}
with ZipFile(REF) as z:
    info['parts']={n:{'size':len(z.read(n)),'sha256':sha256(z.read(n)).hexdigest()} for n in z.namelist()}
(QA/'mau_tham_chieu.json').write_text(json.dumps(info,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'reference':str(REF),'hash':info['sha256'],'paragraphs':len(d.paragraphs),'tables':len(d.tables),'images':len(d.inline_shapes)},ensure_ascii=False))
