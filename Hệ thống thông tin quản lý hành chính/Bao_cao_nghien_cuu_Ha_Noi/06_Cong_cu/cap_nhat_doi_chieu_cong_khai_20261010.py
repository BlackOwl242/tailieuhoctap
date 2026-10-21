from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "04_Bao_cao" / "Phan_tich_thiet_ke_huong_doi_tuong_Ha_Noi.docx"

doc = Document(DOCX)

def find_paragraph(prefix: str):
    return next(p for p in doc.paragraphs if p.text.startswith(prefix))


def add_hyperlink(paragraph, text: str, url: str):
    relationship_id = paragraph.part.relate_to(
        url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(color)
    properties.append(underline)
    run.append(properties)
    label = OxmlElement("w:t")
    label.text = text
    run.append(label)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


find_paragraph("Mỗi phiên bản thủ tục phải có quyết định").text = (
    "Mỗi phiên bản thủ tục phải có quyết định, hiệu lực, cơ quan, địa bàn áp dụng, thành phần, mẫu, "
    "quy tắc hạn, khoản thu, tuyến xử lý và quy tắc ký. Hồ sơ chốt phiên bản và tuyến cơ quan tại thời điểm "
    "gửi hoặc tiếp nhận theo chính sách đã được nghiệp vụ xác nhận. Khi cổng giao dịch cho phép chọn nhiều "
    "địa bàn hoặc loại cơ quan, mã thủ tục không đủ để suy ra một cơ quan duy nhất. Công bố bản mới không "
    "sửa ngầm hồ sơ đang xử lý. Nếu chuyển đổi hồ sơ cũ được phép, phải lưu căn cứ, ánh xạ trạng thái, kết "
    "quả kiểm tra và người phê duyệt."
)

find_paragraph("Bản thiết kế").text = (
    "Bản thiết kế và phần mềm nghiên cứu cục bộ có thể dùng làm cơ sở xây dựng chức năng dùng chung và "
    "thử nghiệm bằng hồ sơ giả lập. Phần mềm React/Python hiện chạy riêng trên máy phát triển; chưa kết nối "
    "dữ liệu, tài khoản, API hoặc giao dịch của hệ thống Hà Nội. Để áp dụng cho từng thủ tục thực tế, cần "
    "bổ sung tờ khai, quy trình hiện hành, quyền của cán bộ và quy định kết nối của đối tác. Danh mục 2.075 "
    "thủ tục mới cung cấp điểm bắt đầu cho việc tra cứu; chưa đủ để đưa toàn bộ danh mục vào vận hành. Việc "
    "triển khai phải thực hiện từng nhóm thủ tục đã được cơ quan phụ trách xác nhận."
)

test_plan = next(
    p for p in doc.paragraphs
    if p.text.startswith("Kế hoạch kiểm thử gồm 168 tình huống")
    or p.text.startswith("Bộ kế hoạch có 96 ca kiểm thử")
    or p.text.startswith("Bộ kế hoạch có 101 ca kiểm thử")
)
test_plan.text = (
    "Bộ kế hoạch có 101 ca kiểm thử gắn với 36 use case, gồm 50 ca trọng tâm, 20 ca mở rộng và 31 ca chức năng. "
    "Các tình huống bao phủ điều kiện hợp lệ, thiếu dữ liệu, sai quyền, gửi lặp, lỗi tích hợp và cập nhật đồng thời. "
    "Đây là kế hoạch; chưa chạy kiểm thử hệ thống với tài khoản, dữ liệu hoặc môi trường do cơ quan vận hành cấp. "
    "Kiểm tra tính nhất quán của mô hình không được tính là kết quả kiểm thử phần mềm."
)

chapter_3 = find_paragraph("CHƯƠNG 3.")
if not any(p.text.startswith("2.4. Đối chiếu cổng công khai") for p in doc.paragraphs):
    heading = chapter_3.insert_paragraph_before(
        "2.4. Đối chiếu cổng công khai", style="Heading 2"
    )
    heading.paragraph_format.keep_with_next = True
    chapter_3.insert_paragraph_before(
        "Trang thông tin điện tử Trung tâm Phục vụ hành chính công Hà Nội công khai các mục Trang chủ, "
        "Giới thiệu, Tin tức sự kiện, Phản ánh kiến nghị, Thư viện số và Liên hệ. Trang chủ Cổng Dịch vụ "
        "công quốc gia có các nhánh Thông tin và dịch vụ, Thanh toán trực tuyến, Phản ánh kiến nghị, Đánh "
        "giá chất lượng phục vụ và Hỗ trợ. Từ 01/07/2025, Hà Nội hướng giao dịch trực tuyến của người dân "
        "về Cổng quốc gia. [2], [11], [12]"
    )
    chapter_3.insert_paragraph_before(
        "Trang đăng nhập công khai phân biệt công dân xác thực bằng VNeID và tổ chức/doanh nghiệp dùng "
        "tài khoản do Cổng quốc gia cấp. Khi người dùng chọn Nộp trực tuyến, hệ thống chuyển tới trang "
        "xác thực. Cổng yêu cầu chọn tỉnh/thành phố và phường/xã hoặc Sở; danh sách phụ thuộc cơ quan đã "
        "áp dụng dịch vụ trực tuyến. Việc kiểm tra một bản ghi xác nhận tuyến: thủ tục 2.000908 tại "
        "Hà Nội → Phường Tây Hồ hiển thị UBND phường Tây Hồ. Đây là quan sát theo một bản ghi và địa bàn, "
        "không được suy rộng sang các địa bàn khác. [13], [14]"
    )
    chapter_3.insert_paragraph_before(
        "Quan sát công khai không làm lộ quy trình sau đăng nhập, quyền cán bộ, API, cấu hình ký số, thanh toán "
        "hoặc đồng bộ. Ứng dụng nghiên cứu giữ riêng bàn cán bộ, nhưng tài khoản, hồ sơ và thao tác ở đây chưa "
        "liên thông hệ thống Hà Nội. Các phần nội bộ được thiết kế theo quy định và vai trò đã công bố, không "
        "được ghi là bản sao cấu hình thật. Cổng quốc gia có quy chế quản lý, vận hành và khai thác riêng. [22]"
    )

if not any(p.text.startswith("2.5. Quy trình chi nhánh") for p in doc.paragraphs):
    heading = chapter_3.insert_paragraph_before(
        "2.5. Quy trình chi nhánh, tiếp nhận và yêu cầu dừng", style="Heading 2"
    )
    heading.paragraph_format.keep_with_next = True
    chapter_3.insert_paragraph_before(
        "Phụ lục Quyết định 2265/QĐ-TTPVHCC công bố các quy trình QT-01 đến QT-11 cho hoạt động tại "
        "chi nhánh và điểm hỗ trợ. QT-03 mô tả việc hướng dẫn người dân tự đăng nhập Cổng quốc gia, "
        "hỗ trợ số hóa, kiểm tra hồ sơ, tiếp nhận hồ sơ điện tử toàn trình hoặc đối chiếu bản giấy với bản điện tử "
        "đối với thủ tục một phần; hồ sơ đủ điều kiện được chuyển cơ quan có thẩm quyền. QT-05 mô tả kiểm tra "
        "định danh, dữ liệu và tính đầy đủ; QT-06 mô tả ký/phát hành, kênh trả, gia hạn và lưu vết. Phụ lục để trống "
        "ngày hiệu lực, vì vậy nó chứng minh nội dung quy trình được công bố nhưng không xác nhận đây là cấu hình "
        "đang chạy của hệ thống. [16]"
    )
    chapter_3.insert_paragraph_before(
        "Quyết định 6470/QĐ-UBND mô tả Trung tâm thực hiện hướng dẫn, hỗ trợ, tiếp nhận, số hóa, phối hợp, "
        "trả kết quả và quản lý hệ thống Thành phố; việc trực tiếp giải quyết chỉ trong phạm vi được giao. "
        "Mô hình do đó tách điểm tiếp nhận khỏi cơ quan chuyên môn có thẩm quyền và người ký. [17]"
    )
    chapter_3.insert_paragraph_before(
        "Nghị định 309/2026 yêu cầu khai thác dữ liệu hợp lệ đã kết nối và hướng dẫn đầy đủ các thiếu sót đầu vào "
        "trong một lần; Nghị định 310/2026 điều chỉnh giao dịch điện tử. Thông tư 13/2026 có hiệu lực ngày 10/10/2026, "
        "bãi bỏ một phần Thông tư 03/2025 và ban hành biểu mẫu 01–08 cùng 9A/9B; Thông tư 14/2026 cùng ngày quy định "
        "số hóa và tái cấu trúc. Hồ sơ tiếp nhận trước ngày hiệu lực phải áp dụng điều khoản chuyển tiếp. [18]–[21]"
    )
    chapter_3.insert_paragraph_before(
        "Đối với yêu cầu trực tiếp hoặc bưu chính, sau khi người dân đã nhận giấy tiếp nhận và trước khi cơ quan ban hành "
        "quyết định hoặc văn bản thông báo kết quả, Bộ phận Một cửa kiểm tra ngay tình trạng hồ sơ. Nếu chưa có kết quả "
        "hoặc văn bản thông báo kết quả thì phát hành Thông báo dừng Mẫu 06; biểu mẫu dành phần ký cho thủ trưởng cơ quan, "
        "đơn vị. Yêu cầu trực tuyến được thực hiện qua chức năng của Cổng quốc gia. Hoàn phí/lệ phí và trả lại giấy tờ "
        "được theo dõi theo quy định của cơ quan thu và quy trình tiếp nhận. Ứng dụng nghiên cứu chỉ mô hình hóa kiểm tra, "
        "lập dự thảo, thẩm quyền, bản ký và chứng cứ giao; chưa ký số hoặc gửi giao dịch lên Cổng. [20]"
    )
    chapter_3.insert_paragraph_before(
        "Đối với thủ tục 2.000635, chỉ mục tìm kiếm của Cổng quốc gia trả về bản ghi 2.000635.H17, cấp xã, "
        "Quyết định 386/QĐ-UBND; trích lục nêu mức 8.000 đồng/bản sao và xử lý trong ngày. Trang chi tiết trực tiếp "
        "trả 503 khi kiểm tra trực tiếp. Chưa xác nhận hiệu lực, khoản thu hoặc cơ quan tại địa bàn Hà Nội; "
        "không dùng thông tin chỉ mục thay cho phiên bản và tuyến hiện hành. [15]"
    )

new_refs = [
    (
        "[11] Trung tâm Phục vụ hành chính công thành phố Hà Nội, Trang thông tin điện tử Trung tâm.",
        "https://ttpvhcc.hanoi.gov.vn/trang-chu",
    ),
    (
        "[12] Cổng Dịch vụ công quốc gia, Trang chủ đang vận hành.",
        "https://dichvucong.gov.vn/",
    ),
    (
        "[13] Cổng Dịch vụ công quốc gia, trang đăng nhập/đăng ký và phương thức xác thực của công dân, tổ chức, doanh nghiệp.",
        "https://dangky.dichvucong.gov.vn/",
    ),
    (
        "[14] Cổng Dịch vụ công quốc gia, bản ghi mã 2.000908 sau khi chọn Thành phố Hà Nội và Phường Tây Hồ.",
        "https://dichvucong.gov.vn/tim-kiem-thu-tuc-hanh-chinh?formalityId=019d2bfd-8e5d-72f9-8cf2-c71780251483&searchType=PROVINCE&province=019bad30-cd82-76ed-b1cc-14070f806e01&provinceCode=01&wardCode=00103&agency=019bad30-cd82-76ed-b1cc-3538098faf4b&commune=WARD&showAdvanced=false&isProvince=0&isMinistry=0",
    ),
    (
        "[15] Cổng Dịch vụ công quốc gia, kết quả chỉ mục tìm kiếm mã 2.000635.H17; trang chi tiết trực tiếp trả 503.",
        "https://thutuc.dichvucong.gov.vn/p/home/dvc-tthc-thu-tuc-hanh-chinh-chi-tiet.html?ma_thu_tuc=182694",
    ),
    (
        "[16] Trung tâm Phục vụ hành chính công thành phố Hà Nội, Quyết định 2265/QĐ-TTPVHCC ngày 29/12/2025, phụ lục QT-01 đến QT-11; trường ngày hiệu lực để trống.",
        "https://datafiles.hanoi.gov.vn/gov-hni/6847/VanBan/2026/1/8/Q%C4%90-TTPVHCC-2265-2025-PL.pdf",
    ),
    (
        "[17] UBND thành phố Hà Nội, Quyết định 6470/QĐ-UBND ngày 26/12/2025 về chức năng, nhiệm vụ, quyền hạn và cơ cấu tổ chức của Trung tâm Phục vụ hành chính công.",
        "https://ttpvhcc.hanoi.gov.vn/chuc-nang-nhiem-vu",
    ),
    (
        "[18] Chính phủ, Nghị định 309/2026/NĐ-CP ngày 05/08/2026 sửa đổi Nghị định 118/2025/NĐ-CP.",
        "https://vanban.chinhphu.vn/?docid=219100&pageid=27160",
    ),
    (
        "[19] Chính phủ, Nghị định 310/2026/NĐ-CP ngày 05/08/2026 sửa đổi quy định thực hiện TTHC trên môi trường điện tử.",
        "https://vanban.chinhphu.vn/?classid=1&docid=219101&pageid=27160",
    ),
    (
        "[20] Bộ Tư pháp, Thông tư 13/2026/TT-BTP, hiệu lực 10/10/2026, quy trình một cửa và bộ biểu mẫu mới.",
        "https://vanban.chinhphu.vn/?docid=219795&pageid=27160",
    ),
    (
        "[21] Bộ Tư pháp, Thông tư 14/2026/TT-BTP, hiệu lực 10/10/2026, số hóa và thực hiện TTHC trên môi trường điện tử.",
        "https://vanban.chinhphu.vn/?docid=219796&pageid=27160&typegroupid=6",
    ),
    (
        "[22] Thủ tướng Chính phủ, Quyết định 29/2026/QĐ-TTg về quản lý, vận hành và khai thác Cổng Dịch vụ công quốc gia.",
        "https://vanban.chinhphu.vn/?docid=218352&orggroupid=3&pageid=27160",
    ),
]
# The first ten bibliography paragraphs were generated with one extra pair of
# brackets. Normalize them so both in-text citations and the audit script can
# resolve the actual reference numbers consistently.
import re
for paragraph in doc.paragraphs:
    match = re.match(r"^\[\[(\d+)\]\](.*)$", paragraph.text)
    if match:
        paragraph.text = f"[{match.group(1)}]{match.group(2)}"

for description, url in new_refs:
    paragraph = next(
        (p for p in doc.paragraphs if p.text.startswith(description)), None
    )
    if paragraph is None:
        appendix = next((p for p in doc.paragraphs if p.text.startswith("PHỤ LỤC.")), None)
        paragraph = appendix.insert_paragraph_before() if appendix is not None else doc.add_paragraph()
        paragraph.style = doc.styles["Normal"]
    linked_text = "".join(node.text or "" for node in paragraph._p.xpath(".//w:hyperlink//w:t"))
    if url not in linked_text:
        for child in list(paragraph._p):
            if child.tag != qn("w:pPr"):
                paragraph._p.remove(child)
        paragraph.add_run(description + " ")
        add_hyperlink(paragraph, url, url)

for paragraph in doc.paragraphs:
    if re.match(r"^\[\d+\]", paragraph.text):
        paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.keep_together = True

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.text = run.text.replace(
                        "168 ca kiểm thử dự kiến", "101 ca kiểm thử dự kiến"
                    )

doc.core_properties.comments = "Đối chiếu nguồn công khai; không xác nhận tích hợp hệ thống nội bộ."
doc.save(DOCX)
print("Updated report DOCX; inserted current public observations and retained the existing report structure.")
