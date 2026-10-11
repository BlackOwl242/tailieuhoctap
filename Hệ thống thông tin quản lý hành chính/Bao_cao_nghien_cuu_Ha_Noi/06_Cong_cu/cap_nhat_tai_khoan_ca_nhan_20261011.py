"""Update current staff screenshots and account documentation without rebuilding UML."""
from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / '04_Bao_cao/Phan_tich_thiet_ke_huong_doi_tuong_Ha_Noi.docx'
UI = ROOT / '03_Thiet_ke/Anh_giao_dien_he_thong'
QA = ROOT / '05_Doi_chieu/qa_staff_account_20261011'
QA.mkdir(parents=True, exist_ok=True)
backup = QA / 'bao_cao_truoc_cap_nhat.docx'
if not backup.exists():
    backup.write_bytes(REPORT.read_bytes())
doc = Document(REPORT)
for p in doc.paragraphs:
    for run in p.runs:
        if '22 bài kiểm thử HTTP' in run.text:
            run.text = run.text.replace('22 bài kiểm thử HTTP', '23 bài kiểm thử HTTP')
        if 'Các màn hình phần mềm nghiên cứu' in run.text:
            run.text = run.text.replace('Các màn hình phần mềm nghiên cứu', 'Các màn hình hệ thống')
files = {
    '4.26': '01_cong_dan_may_tinh.jpg', '4.27': '02_cong_dan_dien_thoai.jpg',
    '4.28': '03_dang_nhap_can_bo.jpg', '4.29': '04_tong_quan_can_bo.jpg',
    '4.30': '05_hop_xu_ly.jpg', '4.31': '06_chi_tiet_ho_so.jpg',
    '4.32': '07_quan_ly_tai_khoan.jpg', '4.33': '08_giam_sat_quy_trinh.jpg',
    '4.34': '09_danh_muc_thu_tuc.jpg', '4.35': '10_nhat_ky_ho_so.jpg',
    '4.36': '11_ket_noi_dich_vu.jpg', '4.38': '13_quan_ly_tin_tuc.jpg',
    '4.37': '12_cong_dan_tin_tuc.jpg',
    '4.39': '14_soan_tin_tuc.jpg',
    '4.40': '15_tai_khoan_ca_nhan.jpg', '4.41': '16_bao_mat_tai_khoan.jpg',
}
changed = []
for paragraph in doc.paragraphs:
    for number, filename in files.items():
        if not paragraph.text.startswith(f'Hình {number}:'):
            continue
        image_node = paragraph._p.getprevious()
        image_paragraph = Paragraph(image_node, doc._body)
        image_paragraph.paragraph_format.keep_with_next = True
        image_paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.keep_together = True
        blips = image_node.xpath('.//a:blip')
        if len(blips) != 1:
            raise RuntimeError(f'Expected one picture before figure {number}')
        image_path = UI / filename
        rel = blips[0].get(qn('r:embed'))
        doc.part.related_parts[rel]._blob = image_path.read_bytes()
        width, height = Image.open(image_path).size
        for extent in image_node.xpath('.//wp:extent'):
            cx = int(extent.get('cx'))
            cy = round(cx * height / width)
            extent.set('cy', str(cy))
            for local in image_node.xpath('.//a:xfrm/a:ext'):
                local.set('cx', str(cx)); local.set('cy', str(cy))
        for props in image_node.xpath('.//wp:docPr'):
            props.set('descr', paragraph.text.split(': ', 1)[1])
        changed.append(number)

heading = next(p for p in doc.paragraphs if p.text.startswith('4.6. Giao diện'))
heading.text = '4.6. Giao diện hệ thống'
intro = heading._p.getnext()
for p in doc.paragraphs:
    if p._p is intro:
        p.text = ('Các ảnh trong mục này được chụp từ Hệ thống thông tin quản lý hành chính, gồm cổng người dân, '
            'các phân hệ cán bộ và phần quản lý tin tức. Bàn cán bộ dùng biểu trưng đã lưu từ cổng thông tin Hà Nội; '
            'tên hệ thống hiển thị trên thanh đầu trang. Nút Đăng xuất nằm ở cuối thanh điều hướng bên trái, '
            'còn tên người dùng trên thanh đầu trang mở mục Tài khoản cá nhân. Trạng thái hồ sơ, số lượng và phạm vi '
            'được trình bày bằng chữ; các chấm báo hoạt động và lối tắt trùng với danh mục điều hướng đã được bỏ. '
            'Các tài khoản và hồ sơ trong ảnh được tạo trong cơ sở dữ liệu riêng để minh họa thao tác.')

caption = next(p for p in doc.paragraphs if p.text.startswith('Hình 4.26:'))
anchor = caption._p.getprevious()
text = ('Tài khoản cá nhân hiển thị họ tên, email đăng nhập, vai trò, phạm vi cơ quan, thời hạn hiệu lực, '
    'quyền ký Mẫu 06 và căn cứ phê duyệt nếu có. Người dùng cập nhật số điện thoại sau khi nhập đúng mật khẩu hiện tại; '
    'họ tên, email đăng nhập và quyền công tác vẫn được điều chỉnh qua quản trị viên. Chức năng đổi mật khẩu áp dụng '
    'cho tài khoản đang sử dụng, yêu cầu từ 12 ký tự gồm chữ hoa, chữ thường và chữ số. Khi đổi mật khẩu, máy chủ '
    'kết thúc các phiên khác và giữ phiên đang thao tác. Mục Phiên đăng nhập cho biết số phiên còn hiệu lực, hạn '
    'của phiên hiện tại và cho phép kết thúc các phiên khác. Phê duyệt cấp quyền và xác minh quyền sở hữu email '
    'là hai việc riêng: hệ thống hiện chưa gửi mã hoặc liên kết xác minh email, nên giao diện ghi rõ Chưa xác minh '
    'qua email. Cờ xác minh cũ trong dữ liệu không được sử dụng làm bằng chứng email đã được kiểm tra.')
if not any(p.text.startswith('Tài khoản cá nhân hiển thị') for p in doc.paragraphs):
    p = doc.add_paragraph(text)
    anchor.addprevious(p._p)

for table in doc.tables:
    if table.rows[0].cells[0].text == 'Khu vực' and table.rows[0].cells[1].text == 'Màn hình thể hiện':
        if not any(row.cells[0].text == 'Tài khoản cá nhân' for row in table.rows):
            row = table.add_row()
            row.cells[0].text = 'Tài khoản cá nhân'
            row.cells[1].text = 'Thông tin quyền và xác minh; cập nhật liên hệ; đổi mật khẩu; quản lý các phiên đăng nhập'
        for row in table.rows:
            if row.cells[0].text == 'Tài khoản cá nhân':
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        p.paragraph_format.first_line_indent = Pt(0)

last_caption = next(p for p in doc.paragraphs if p.text.startswith('Hình 4.39:'))
chapter5 = next(p for p in doc.paragraphs if p.text.startswith('CHƯƠNG 5.'))
for number, filename, title in [
    (40, '15_tai_khoan_ca_nhan.jpg', 'Tài khoản cá nhân, quyền công tác và tình trạng xác minh email'),
    (41, '16_bao_mat_tai_khoan.jpg', 'Cập nhật liên hệ, đổi mật khẩu và quản lý phiên đăng nhập'),
]:
    if any(p.text.startswith(f'Hình 4.{number}:') for p in doc.paragraphs):
        continue
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = True
    picture = p.add_run().add_picture(str(UI / filename), width=Inches(6.3))
    picture._inline.docPr.set('descr', title)
    chapter5._p.addprevious(p._p)
    copy = deepcopy(last_caption._p)
    nodes = copy.xpath('.//w:t')
    nodes[0].text = 'Hình 4.'
    nodes[1].text = str(number)
    nodes[-1].text = ': ' + title
    chapter5._p.addprevious(copy)

doc.save(REPORT)
print('Updated figures:', ', '.join(changed), '; added 4.40 and 4.41; saved', REPORT)
