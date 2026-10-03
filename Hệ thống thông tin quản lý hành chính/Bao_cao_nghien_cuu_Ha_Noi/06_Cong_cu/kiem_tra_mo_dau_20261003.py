from pathlib import Path
from hashlib import sha256
from zipfile import ZipFile
import json, re
from docx import Document
import ket_xuat_va_doi_chieu_oo as qa

ROOT = Path(__file__).resolve().parents[1]
Q = ROOT / '05_Doi_chieu/Mo_dau_20261003'
DOC = ROOT / '04_Bao_cao/Phan_tich_thiet_ke_huong_doi_tuong_Ha_Noi.docx'
qa.Q = Q
pages = qa.render(Q / 'Bao_cao.pdf', Q / 'Trang')
qa.check(DOC)
before = json.loads((Q / 'Trang_truoc_sua.json').read_text('utf-8'))
changed = [p.name for p in sorted((Q / 'Trang').glob('trang_*.png')) if sha256(p.read_bytes()).hexdigest() != before.get(p.name)]
old_text = json.loads((ROOT / '05_Doi_chieu/Bien_tap_lai_20261003/Trang/van_ban_theo_trang.json').read_text('utf-8'))
text_changes = [b['page'] for a, b in zip(old_text, pages) if a['text'] != b['text']]
old, new = Document(Q / 'Ban_truoc_sua.docx'), Document(DOC)

def after_opening(d):
    paragraphs = [p.text for p in d.paragraphs]
    start = paragraphs.index('MỤC LỤC')
    return paragraphs[start:], [[c.text for row in t.rows for c in row.cells] for t in d.tables]

text = '\n'.join(p.text for p in new.paragraphs) + '\n' + '\n'.join(c.text for t in new.tables for row in t.rows for c in row.cells)
specs = [t for t in new.tables if len(t.rows) > 1 and t.cell(1, 0).text == 'Mã và mục tiêu']
with ZipFile(Q / 'Ban_truoc_sua.docx') as zold, ZipFile(DOC) as znew:
    media = [n for n in zold.namelist() if n.startswith('word/media/')]
    media_unchanged = all(zold.read(n) == znew.read(n) for n in media)
checks = {
    'Tổng số trang giữ nguyên': len(pages) == len(before) == 206,
    'Chỉ nội dung hai trang mở đầu thay đổi': text_changes == [3, 4],
    'Nội dung từ mục lục trở đi giữ nguyên': after_opening(old) == after_opening(new),
    'Toàn bộ hình ảnh giữ nguyên': media_unchanged,
    'Đủ 36 đặc tả Use case': len(specs) == 36,
    'Mỗi Use case có tác nhân': all(any('Tác nhân' in row.cells[0].text for row in t.rows) for t in specs),
    'Không có mã UCN hoặc R': not re.search(r'\b(?:UCN|R)\d+\b', text),
    'Không nhắc tài liệu mẫu trong báo cáo': not any(s in text.lower() for s in ['pttk_oop_hr', 'talentconnect', 'tài liệu mẫu', 'audit mẫu']),
    'Đủ 10 nguồn trích dẫn': set(re.findall(r'\[(\d+)\]', text)) == set(map(str, range(1, 11))),
    'Không có chữ vượt trang': all(not p['overflow'] for p in pages),
}
result = {'docx': str(DOC), 'sha256': sha256(DOC.read_bytes()).hexdigest(), 'pages': len(pages), 'changed_pages': changed, 'text_changes': text_changes, 'unchanged_pages': len(pages)-len(changed), 'media_count': len(media), 'checks': checks, 'scope': 'Chỉ sửa lời cảm ơn và lời cam đoan. Phần tác nhân và thiết kế giữ nguyên; không xác nhận phần mềm thực tế đã được kiểm thử.'}
(Q / 'Kiem_tra_mo_dau.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), 'utf-8')
print(json.dumps(result, ensure_ascii=False, indent=2))
