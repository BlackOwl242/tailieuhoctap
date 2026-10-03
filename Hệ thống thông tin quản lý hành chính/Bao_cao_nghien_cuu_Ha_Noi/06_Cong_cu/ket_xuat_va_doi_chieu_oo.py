from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from hashlib import sha256
import json,re,sys
from docx import Document
import pypdfium2 as pdfium
import pdfplumber
ROOT=Path(__file__).resolve().parents[1];Q=ROOT/'05_Doi_chieu/Kiem_tra_huong_doi_tuong_20261002';REF=ROOT.parents[1]/'Công nghệ phần mềm/Bao_Cao_Bai_Tap_Lon_CNPM_TalentConnect.docx'
def restore(file):
 with ZipFile(REF) as z,ZipFile(file) as f:
  data={n:f.read(n) for n in f.namelist()}
  for n in z.namelist():
   if n.startswith(('word/theme/','word/footer','word/header','word/numbering','word/styles','word/fontTable')):data[n]=z.read(n)
 tmp=file.with_suffix('.preserved.docx')
 with ZipFile(tmp,'w',ZIP_DEFLATED) as z:
  for n,b in data.items():z.writestr(n,b)
 tmp.replace(file)
def render(pdf,out):
 out.mkdir(exist_ok=True);d=pdfium.PdfDocument(str(pdf));texts=[]
 with pdfplumber.open(pdf) as p:
  for i,page in enumerate(d):
   page.render(scale=1.5).to_pil().save(out/f'trang_{i+1:03}.png');words=p.pages[i].extract_words();texts.append({'page':i+1,'text':p.pages[i].extract_text() or '', 'words':len(words),'overflow':[w['text'] for w in words if w['x0']<-1 or w['x1']>p.pages[i].width+1 or w['top']<-1 or w['bottom']>p.pages[i].height+1]})
 (out/'van_ban_theo_trang.json').write_text(json.dumps(texts,ensure_ascii=False,indent=2),encoding='utf-8');print('render',pdf.name,len(d));return texts
def check(file):
 d=Document(file);ref=Document(REF);checks=[]
 def ck(t,b):checks.append({'check':t,'passed':bool(b)})
 ck('Hai phần theo mẫu',len(d.sections)==len(ref.sections)==2)
 for i,(s,r) in enumerate(zip(d.sections,ref.sections)):
  for k in ['page_width','page_height','top_margin','bottom_margin','left_margin','right_margin','header_distance','footer_distance']:ck(f'Phần {i+1}, {k}',getattr(s,k)==getattr(r,k))
 with ZipFile(REF) as z,ZipFile(file) as f:
  parts=[n for n in z.namelist() if n.startswith(('word/theme/','word/footer','word/header','word/numbering','word/styles','word/fontTable'))]
  for n in parts:ck('Bảo toàn '+n,z.read(n)==f.read(n))
 text='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells)
 ck('Tên sinh viên đúng', 'Lê Quóc Huy' in text and '2305HTTB011' in text)
 ck('Giảng viên đúng','Hoàng Minh Ngọc' in text)
 ck('Không có dấu gạch kép','--' not in text)
 ck('Không có thuật ngữ tiếng Anh thêm trong ngoặc',not re.search(r'\((?:workflow|use case|class|sequence|activity|deployment|component|business|API|UML)[^)]*\)',text,re.I))
 ck('Không còn chuỗi lỗi trường Word','Error!' not in text and 'Lỗi!' not in text)
 caps=[p.text for p in d.paragraphs if re.match(r'^Hình [0-9A-Z]+\.\d+:',p.text) and '\t' not in p.text]
 ck('Mỗi hình có chú thích không tính danh mục',len(caps)==(160 if 'Phan_tich' in file.name else 0))
 ck('Bảng không chia hàng',all(row._tr.find(qn('w:trPr')).find(qn('w:cantSplit')) is not None for t in d.tables for row in t.rows))
 result={'docx':str(file),'sha256':sha256(file.read_bytes()).hexdigest(),'checks':checks,'all_passed':all(x['passed'] for x in checks),'true_figures':len(caps),'tables':len(d.tables),'meaning':'Chỉ xác nhận các tiêu chí cấu trúc được liệt kê; không khẳng định mọi yêu cầu nghiệp vụ đã đạt 100%'}
 (Q/(file.stem+'_kiem_tra.json')).write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');print(file.name,'checks',len(checks),'failed',[x['check'] for x in checks if not x['passed']])
from docx.oxml.ns import qn
if __name__=='__main__':
 mode=sys.argv[1]
 if mode=='restore':
  for f in (ROOT/'04_Bao_cao').glob('*.docx'):
   if f.name.startswith(('Phan_tich','Tom_tat')):restore(f)
 elif mode=='render':
  for pdf in [Q/'Thiet_ke_cuoi.pdf',Q/'Tom_tat_cuoi.pdf']:render(pdf,Q/pdf.stem)
 elif mode=='check':
  for f in (ROOT/'04_Bao_cao').glob('*.docx'):
   if f.name.startswith(('Phan_tich','Tom_tat')):check(f)
