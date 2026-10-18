"""Các đoạn thuyết minh dùng khi biên tập báo cáo học phần."""
import re
from copy import deepcopy

HEADINGS = {
 '1.2. Đơn vị đếm chức năng và quy trình':'1.2. Phạm vi chức năng và quy trình',
 '1.3. Phương pháp và mức độ chứng minh':'1.3. Phương pháp nghiên cứu',
 'CHƯƠNG 3. TÁC NHÂN, Use case VÀ HƯỚNG DẪN CHỨC NĂNG':'CHƯƠNG 3. TÁC NHÂN VÀ CÁC CA SỬ DỤNG',
 '3.3. Đặc tả 36 Use case theo mã':'3.3. Đặc tả các ca sử dụng',
 '6.4. Mức độ sẵn sàng và dữ liệu bàn giao':'6.4. Kết quả triển khai và hạn chế',
 '6.5. Đối chiếu Use case với bản triển khai hiện tại':'6.5. Đánh giá mức độ đáp ứng các ca sử dụng',
}

REVISIONS = [
 ('Trong quá trình thực hiện, em đã đối chiếu các nhánh dịch vụ công khai', 'Trong quá trình thực hiện đề tài, em đã tìm hiểu các dịch vụ công khai, phân tích 36 ca sử dụng và xây dựng mô hình dữ liệu cho hệ thống. Từ kết quả phân tích, em phát triển ứng dụng quản lý hồ sơ từ tiếp nhận đến giao kết quả, đồng thời xây dựng hướng dẫn sử dụng cho từng nhóm người dùng.'),
 ('Em mong nhận được nhận xét của giảng viên về căn cứ', 'Em mong nhận được ý kiến góp ý của giảng viên để tiếp tục hoàn thiện nội dung phân tích, thiết kế và ứng dụng của đề tài.'),
 ('Em cam đoan báo cáo do em thực hiện', 'Em cam đoan báo cáo là kết quả nghiên cứu và thực hiện của em trong học phần Hệ thống thông tin quản lý hành chính. Các tài liệu, số liệu và hình ảnh tham khảo đều được ghi nguồn. Em chịu trách nhiệm về nội dung phân tích, thiết kế và kết quả trình bày trong báo cáo.'),
 ('Ứng dụng React/Python trong dự án vận hành độc lập', 'Ứng dụng được xây dựng và vận hành trong phạm vi đề tài học phần. Các hình ở mục 4.6 được chụp từ ứng dụng này; hình giao diện công khai và tư liệu lịch sử được ghi nguồn ở chương 2.'),
 ('Em chưa được cấp quyền khảo sát tài khoản cán bộ', 'Việc khảo sát dựa trên tài liệu và thông tin công khai. Những nội dung về cấu hình nội bộ và kết nối với hệ thống của Thành phố được trình bày trong phần hạn chế và hướng phát triển.'),
 ('Đối tượng nghiên cứu là Hệ thống thông tin giải quyết thủ tục hành chính thành phố Hà Nội.', 'Đề tài nghiên cứu Hệ thống thông tin giải quyết thủ tục hành chính thành phố Hà Nội, được ghi nhận đưa vào vận hành từ ngày 11/04/2023 [10]. Tài liệu hướng dẫn công dân [1] mô tả các chức năng của hệ thống tại thời điểm công bố. Từ ngày 01/07/2025, người dân được hướng dẫn thực hiện thủ tục qua Cổng Dịch vụ công quốc gia [2]. Trong báo cáo, cổng tiếp nhận dành cho người dân và hệ thống xử lý hồ sơ của Thành phố được xem xét theo từng chức năng để làm rõ trách nhiệm của mỗi bên.'),
 ('Một chức năng có thể phục vụ hàng nghìn thủ tục.', 'Chức năng phần mềm, thủ tục hành chính và quy trình xử lý có phạm vi khác nhau. Một chức năng như tiếp nhận hay tra cứu có thể phục vụ nhiều thủ tục, trong khi mỗi thủ tục có thể phát sinh nhiều nhánh tùy chủ thể, cơ quan và hình thức nộp. Vì vậy, báo cáo trình bày riêng số nhóm chức năng, số thủ tục tham chiếu và các nhánh quy trình được phân tích.'),
 ('Nguồn được phân biệt theo loại bằng chứng.', 'Đề tài sử dụng ba nhóm nguồn: giao diện dịch vụ công khai, tài liệu hướng dẫn và văn bản quy định quy trình nghiệp vụ. Mỗi nhóm cung cấp thông tin ở một phạm vi khác nhau. Các nguồn được đối chiếu để xác định yêu cầu, sau đó chuyển thành mô hình phân tích và thiết kế. Tính nhất quán của mô hình được kiểm tra thông qua quan hệ dữ liệu, ca sử dụng và các điều kiện chuyển trạng thái.'),
 ('Ký pháp được đối chiếu với [8].', 'Các sơ đồ sử dụng ký pháp UML theo tài liệu tham chiếu [8]. Biểu đồ trình tự mô tả thứ tự trao đổi giữa các đối tượng; biểu đồ hoạt động thể hiện các bước, điều kiện rẽ nhánh và công việc song song. Biểu đồ lớp mô tả thuộc tính, thao tác và quan hệ dữ liệu. Quy ước về lớp giao diện, lớp điều khiển và lớp thực thể được giải thích ở chương 4.'),
 ('Các ảnh trên dừng ở phần công khai;', 'Các hình ở mục này minh họa việc tra cứu và lựa chọn thủ tục trên cổng công khai. Khi nộp hồ sơ, người dùng lựa chọn cơ quan giải quyết để xem thông tin áp dụng tại địa phương. Thông tin trên cổng quốc gia được đối chiếu với quy trình riêng của Hà Nội khi xác định thành phần, thời hạn và thẩm quyền.'),
 ('Tất cả ảnh của mục 2.2 là tư liệu lịch sử.', 'Các hình ở mục 2.2 được trích từ tài liệu hướng dẫn năm 2023 và dùng để tìm hiểu cách tổ chức giao diện ở giai đoạn đó. Tên cơ quan và dữ liệu trong ảnh được giữ theo tài liệu nguồn. Hướng dẫn thao tác với ứng dụng của đề tài được trình bày riêng ở mục 4.6 và phần cẩm nang sử dụng.'),
 ('Tác nhân là vai trò tương tác, không đồng nhất với tên cơ quan.', 'Tác nhân được xác định theo vai trò tham gia xử lý hồ sơ. Một người có thể được giao nhiều vai trò, nhưng quyền thao tác phụ thuộc nhiệm vụ, cơ quan và thời hạn phân công. Mô hình thiết kế phân biệt quyền quản trị kỹ thuật với quyền nghiệp vụ. Trong ứng dụng hiện tại, quản trị viên được cấp cả quyền quản lý tài khoản và xem dữ liệu hồ sơ; việc tách hai phạm vi này là nội dung cần phát triển thêm.'),
 ('Biểu đồ trình tự của mỗi ca thể hiện một kịch bản tại một thời điểm', 'Mỗi biểu đồ trình tự mô tả một kịch bản cụ thể, với các đối tượng và thông điệp được sắp theo thứ tự xử lý. Các ký hiệu giao diện, điều khiển và thực thể giúp nhận diện trách nhiệm của từng đối tượng. Đối với yêu cầu bổ sung và phát hành kết quả, thao tác của người dân, chuyên viên, lãnh đạo và văn thư được tách rõ theo vai trò. Trong quy trình phối hợp, cơ quan chủ trì chỉ tổng hợp khi đã nhận đủ ý kiến hợp lệ; trường hợp chậm phản hồi được theo dõi để xử lý theo thẩm quyền.'),
 ('Lớp giao diện tiếp nhận dữ liệu và thể hiện lỗi', 'Mô hình phân tích gồm lớp giao diện, lớp điều khiển và lớp thực thể. Lớp giao diện nhận thông tin và hiển thị kết quả cho người dùng. Lớp điều khiển kiểm tra quyền, tổ chức thao tác và quản lý giao dịch. Các thực thể lưu dữ liệu hồ sơ, kết quả, khoản thu và nhiệm vụ phối hợp, đồng thời duy trì những ràng buộc nghiệp vụ. Việc trao đổi với dịch vụ ngoài được thực hiện qua lớp kết nối để thuận tiện thay đổi nhà cung cấp.'),
 ('Các hình lớp trình bày một số thuộc tính và quan hệ chính', 'Các biểu đồ lớp trình bày những thuộc tính và quan hệ chính; bảng 4.1 liệt kê 74 khóa ngoại của mô hình dữ liệu. Bội số ở hai đầu quan hệ cho biết số đối tượng có thể liên kết với nhau. Quan hệ lưu trữ được thiết kế để giữ hồ sơ, tài liệu đã ký và nhật ký khi các thông tin danh mục hoặc tài khoản thay đổi.'),
 ('Từ điển 37 bảng với đầy đủ kiểu, khóa và tính bắt buộc', 'Từ điển dữ liệu mô tả 37 bảng cùng kiểu dữ liệu, khóa và các trường bắt buộc. Các cấu trúc tờ khai được quản lý theo phiên bản của từng thủ tục. Bên cạnh ràng buộc khóa ngoại, lớp ứng dụng kiểm tra quyền truy cập, điều kiện xử lý và vòng đời của đối tượng để bảo đảm dữ liệu được sử dụng đúng nghiệp vụ.'),
 ('Bốn trạng thái kết thúc không có lệnh chuyển thông thường.', 'Bốn trạng thái kết thúc khép lại quy trình xử lý thông thường. Nếu phát sinh đính chính, hủy văn bản hoặc khiếu nại, hệ thống cần lưu thành nghiệp vụ riêng có căn cứ và người chịu trách nhiệm. Hạn giải quyết được tính từ phiên bản quy tắc và lịch làm việc; tình trạng quá hạn được theo dõi như một chỉ tiêu bổ sung. Mọi điều chỉnh thời hạn cần có căn cứ và lịch sử thay đổi.'),
 ('Gói tổ chức mã theo nghiệp vụ;', 'Kiến trúc được trình bày ở ba góc nhìn: các gói nghiệp vụ, các thành phần phần mềm và cách bố trí thành phần trên hạ tầng. Sơ đồ cộng tác bổ sung thứ tự trao đổi khi phê duyệt và phát hành kết quả. Thành phần kết nối sử dụng hợp đồng giao tiếp của nghiệp vụ, nhờ đó có thể thay dịch vụ ký hoặc thanh toán mà vẫn giữ các quy tắc xử lý hồ sơ.'),
 ('Lõi miền và hợp đồng chung không phụ thuộc lớp giao diện', 'Lõi nghiệp vụ quản lý các quy tắc dùng chung và cung cấp giao diện cho những gói chức năng khác. Gói thống kê sử dụng dữ liệu đọc để tổng hợp báo cáo; gói kết nối trao đổi với đối tác thông qua hợp đồng đã xác định. Ma trận phụ thuộc thể hiện đầy đủ các quan hệ này, còn sơ đồ chỉ giữ những đường liên kết chính để thuận tiện theo dõi.'),
 ('Thư viện tại 07_Huong_dan/Thu_tuc_theo_QD_492', 'Thư viện hướng dẫn được xây dựng từ 2.075 dòng thủ tục trong ba phụ lục của nguồn [4]. Người dùng có thể tìm theo tên, mã, cơ quan hoặc lĩnh vực trên trang tra cứu đi kèm. Mỗi mục ghi mức cung cấp dịch vụ trực tuyến, nguồn tham khảo và hướng dẫn tiếp cận. Việc xác định thành phần hồ sơ và điều kiện áp dụng cụ thể tiếp tục dựa vào văn bản còn hiệu lực của từng thủ tục.'),
 ('Một số mã trong bản phụ lục có cách ghi không theo khuôn phổ biến.', 'Thư viện giữ nguyên mã thủ tục theo phụ lục nguồn và đánh dấu những trường hợp cần kiểm tra. Mỗi dòng có một mã quản lý riêng để tránh nhầm lẫn khi số thứ tự trong phụ lục bị trùng. Các nội dung chưa đủ tài liệu chi tiết, như thành phần, số lượng, mẫu, phí hoặc thời hạn, được ghi rõ để tiếp tục đối chiếu.'),
 ('Mã thủ tục là 2.000635.', 'Thủ tục cấp bản sao trích lục hộ tịch, bản sao giấy khai sinh có mã 2.000635. Báo cáo đối chiếu thông tin công khai [7] với phương án tái cấu trúc ngày 24/12/2025 [5]. Việc lựa chọn Hà Nội và cơ quan tiếp nhận là cơ sở xác định quy trình áp dụng; các nội dung về hồ sơ, khoản thu và thời hạn được kiểm tra theo từng nguồn.'),
 ('Mã thủ tục là 2.000908.', 'Thủ tục cấp bản sao từ sổ gốc có mã 2.000908. Quy trình tại cấp xã được công bố ở trang 65–69 của nguồn [6], gồm tiếp nhận trực tiếp, trực tuyến và qua bưu chính, với mức thu bằng không. Hạn giải quyết được xác định theo thời điểm tiếp nhận, kể cả trường hợp sau 15 giờ. Thời gian phân bổ cho từng bước nội bộ được theo dõi riêng để bảo đảm hoàn thành trong hạn.'),
 ('Hai thủ tục trên là trường hợp đối chiếu chuyên sâu', 'Hai thủ tục trên được lựa chọn để phân tích sâu vì có văn bản mô tả quy trình và nguồn công khai để đối chiếu. Kết quả phân tích là cơ sở xây dựng cách cấu hình biểu mẫu, thời hạn và trách nhiệm xử lý. Với các thủ tục khác, cần tiếp tục kiểm tra văn bản thay thế, cơ quan có thẩm quyền và ngày áp dụng trước khi cấu hình.'),
 ('Hợp đồng 36 thao tác được lưu tại', 'Bộ hợp đồng giao tiếp mô tả 36 thao tác tương ứng với các ca sử dụng. Mỗi hợp đồng xác định dữ liệu đầu vào, đầu ra, điều kiện và lỗi có thể phát sinh. Các thao tác đọc sử dụng GET; các lệnh thay đổi sử dụng POST kèm khóa chống lặp. Máy chủ kiểm tra vai trò và phạm vi cơ quan trước khi trả dữ liệu hoặc thực hiện thay đổi. Đặc tả chi tiết được lưu trong bộ thiết kế đi kèm.'),
 ('Trường đối tượng và danh sách cấu trúc phải có lược đồ', 'Các trường có cấu trúc được mô tả chi tiết theo từng thủ tục để kiểm tra kiểu dữ liệu, tài liệu bắt buộc và dung lượng tệp. Khi tích hợp với đối tác, hợp đồng cần được bổ sung bằng lược đồ và quy tắc xác thực do đơn vị quản lý cung cấp. Kết quả gọi thực tế được kiểm tra trên môi trường kết nối trước khi đưa vào sử dụng.'),
 ('UC13 và UC16 là mục tiêu nghiệp vụ nhiều vai trò.', 'UC13 và UC16 có sự tham gia của nhiều vai trò. Người nộp cung cấp tài liệu bổ sung, còn cán bộ kiểm tra và quyết định việc tiếp tục xử lý. Trong quy trình phát hành, lãnh đạo phê duyệt dự thảo, người có quyền ký xác nhận nội dung và văn thư cấp số phát hành. Chín lệnh trong bảng 6.37 mô tả chi tiết trách nhiệm của các vai trò trong hai ca sử dụng này.'),
 ('Bước 1 xác thực người hoặc dịch vụ gọi.', 'Khi nhận yêu cầu, máy chủ xác thực người dùng, xác định phạm vi cơ quan và kiểm tra quyền đối với hồ sơ. Tiếp theo, hệ thống kiểm tra khóa chống lặp, phiên bản dữ liệu và điều kiện chuyển bước. Nếu yêu cầu hợp lệ, thay đổi nghiệp vụ, nhật ký và thông điệp chờ gửi được lưu trong cùng một giao dịch. Kết quả chỉ được trả về sau khi lưu thành công. Việc gửi thông điệp cho đối tác được xử lý sau đó để thời gian phản hồi của dịch vụ ngoài không kéo dài giao dịch hồ sơ.'),
 ('Mỗi phiên bản thủ tục phải có quyết định, hiệu lực', 'Mỗi phiên bản thủ tục lưu căn cứ ban hành, thời gian hiệu lực, cơ quan, thành phần hồ sơ, biểu mẫu, thời hạn, khoản thu và tuyến xử lý. Hồ sơ được gắn với phiên bản áp dụng tại thời điểm gửi hoặc tiếp nhận. Khi công bố phiên bản mới, các hồ sơ đang xử lý vẫn giữ thông tin đã chốt. Nếu cần chuyển đổi, hệ thống lưu căn cứ, người phê duyệt và kết quả đối chiếu.'),
 ('Bộ kiểm thử hướng đối tượng gồm 180 ca dự kiến', 'Kế hoạch kiểm thử gồm 180 tình huống, được liên kết với 36 ca sử dụng để kiểm tra luồng thành công, điều kiện từ chối, phạm vi truy cập, gửi lặp và xung đột dữ liệu. Bộ kiểm tra mô hình đã thực hiện 11 phép kiểm tra về tính duy nhất, quan hệ tham chiếu, khóa ngoại và khả năng đi tới các trạng thái. Các ca kiểm thử phần mềm và kết nối cần tiếp tục được thực hiện, ghi nhận kết quả theo môi trường triển khai.'),
 ('Giám sát tách thời gian phản hồi, lỗi kỹ thuật', 'Hoạt động giám sát theo dõi đồng thời chất lượng dịch vụ và tiến độ nghiệp vụ. Các chỉ tiêu gồm thời gian phản hồi, lỗi kỹ thuật, thông điệp chờ gửi, hồ sơ đến hạn, ý kiến phối hợp và chênh lệch tài chính. Khi đối tác gặp sự cố, đầu mối xử lý lưu thông tin liên quan, theo dõi hồ sơ bị ảnh hưởng và tiếp tục phục vụ theo phương án đã được phê duyệt.'),
 ('Chủ quản nghiệp vụ phê duyệt danh mục và quy trình;', 'Việc vận hành cần có sự phối hợp giữa chủ quản nghiệp vụ, chủ quản dữ liệu, đơn vị kỹ thuật và cán bộ lưu trữ. Mỗi bên được giao trách nhiệm cụ thể đối với cấu hình, quyền khai thác, sao lưu, khôi phục và bảo quản hồ sơ. Những thay đổi ảnh hưởng đến quy trình hoặc dữ liệu cần được phê duyệt, ghi nhận và kiểm tra trước khi áp dụng.'),
 ('Ứng dụng vận hành đầu-cuối các quy trình nội bộ nêu trên;', 'Hướng dẫn khởi chạy được lưu trong README của thư mục dự án. Dữ liệu hồ sơ, tài liệu và tài khoản được lưu để tiếp tục xử lý giữa các phiên làm việc. Cẩm nang sử dụng và bảng đánh giá chức năng đi kèm giúp đối chiếu thao tác trên giao diện với nội dung phân tích trong báo cáo.'),
]

UC_NARRATIVES = {
 'UC01':'Việc cấp và thu hồi quyền được thực hiện theo đề nghị có căn cứ. Sau khi người có thẩm quyền phê duyệt, hệ thống áp dụng quyền mới và lưu lịch sử để xác định ai được thao tác trong từng thời điểm.',
 'UC02':'Danh mục cung cấp mã và thông tin thống nhất cho các phân hệ. Mỗi lần cập nhật cần gắn với nguồn công bố và thời gian hiệu lực, đồng thời giữ thông tin cũ để giải thích các hồ sơ đã tiếp nhận.',
 'UC03':'Phiên bản thủ tục tập hợp thành phần hồ sơ, thời hạn và căn cứ áp dụng. Hồ sơ mới sử dụng phiên bản đã công bố theo ngày hiệu lực; hồ sơ đang giải quyết tiếp tục giữ phiên bản đã tiếp nhận.',
 'UC04':'Cấu hình trạng thái xác định những bước được phép thực hiện và điều kiện đi kèm. Bản cấu hình được kiểm tra đường đi và trình duyệt trước khi áp dụng, trong khi khoản thu và giao nhận vẫn được theo dõi riêng.',
 'UC05':'Người nộp lựa chọn thủ tục và cơ quan, kê khai thông tin, đính kèm tài liệu rồi xem lại trước khi gửi. Hệ thống kiểm tra dữ liệu, cấp mã giao dịch và xử lý yêu cầu gửi lặp theo cùng mã để tránh tạo thêm hồ sơ.',
 'UC06':'Việc luân chuyển xác định rõ đơn vị giao, đơn vị nhận và trách nhiệm xử lý. Cơ quan nguồn theo dõi xác nhận của bên nhận, đồng thời giữ mã hồ sơ và thời hạn để bảo đảm công việc được tiếp tục đầy đủ.',
 'UC07':'Cán bộ mở công việc được giao, xem thông tin và ghi nhận kết quả xử lý ở bước hiện tại. Hệ thống kiểm tra quyền và phiên bản dữ liệu trước khi lưu, giúp tránh ghi đè khi nhiều người cùng làm việc.',
 'UC08':'Tiến trình hồ sơ được trình bày theo các mốc tiếp nhận, xử lý và trả kết quả. Người truy cập chỉ xem nội dung thuộc phạm vi được phép và có thể nhận biết những việc cần thực hiện tiếp theo.',
 'UC09':'Bản quy trình được thiết kế bằng các bước, điều kiện và mối liên hệ giữa chúng. Hệ thống kiểm tra cấu trúc trước khi trình duyệt, đặc biệt với các nhánh song song cần có quy tắc tổng hợp kết quả.',
 'UC10':'Người ký kiểm tra đúng bản dự thảo đã duyệt trước khi xác nhận ký. Kết quả kiểm tra chữ ký được lưu cùng tài liệu; nếu nội dung thay đổi, bản mới cần được duyệt và ký lại theo thẩm quyền.',
 'UC11':'Dịch vụ ký được cấu hình và kiểm tra với tài liệu được phép sử dụng trước khi kết nối nghiệp vụ. Thông tin bí mật được quản lý riêng, còn kết quả ký chỉ được ghi nhận sau khi hệ thống kiểm tra phản hồi hợp lệ.',
 'UC12':'Cán bộ đối chiếu chủ thể, thành phần hồ sơ và thẩm quyền để quyết định tiếp nhận, yêu cầu bổ sung hoặc từ chối. Với hồ sơ đủ điều kiện, hệ thống lập giấy tiếp nhận và tính hạn theo quy tắc của thủ tục.',
 'UC13':'Yêu cầu bổ sung ghi rõ tài liệu còn thiếu và căn cứ để người nộp thực hiện. Sau khi nhận tài liệu, cán bộ kiểm tra và tiếp tục xử lý từ giai đoạn đã lưu; các thay đổi về thời hạn được ghi nhận theo quy định áp dụng.',
 'UC14':'Cán bộ được phân công phải phù hợp với cơ quan và trách nhiệm giải quyết. Khi giao lại hồ sơ, hệ thống kết thúc phân công cũ và lưu phân công mới trong cùng giao dịch để xác định người chịu trách nhiệm.',
 'UC15':'Chuyên viên đối chiếu thông tin, thu thập ý kiến phối hợp khi cần và chuẩn bị dự thảo. Ý kiến được quản lý theo phiên bản để lãnh đạo biết nội dung nào đã được xác nhận khi xem xét kết quả.',
 'UC16':'Lãnh đạo xem dự thảo và căn cứ để phê duyệt hoặc yêu cầu sửa. Sau bước ký, văn thư kiểm tra điều kiện phát hành, cấp số và lưu bản kết quả đã chốt để phục vụ việc giao nhận và tra cứu.',
 'UC17':'Bộ phận Một cửa tiếp nhận yêu cầu Mẫu 05 và kiểm tra chủ thể, thời điểm cùng tình trạng giải quyết. Nếu đủ điều kiện, Mẫu 06 được lập để người có thẩm quyền ký, phát hành rồi mới ghi nhận dừng. Việc trả giấy tờ và hoàn phí được theo dõi theo từng nghĩa vụ.',
 'UC18':'Trước khi giao kết quả, cán bộ xác minh người nhận và phương thức đã đăng ký. Thông tin giao thành công hoặc lý do chưa giao được được lưu cùng chứng cứ để tiếp tục theo dõi trách nhiệm trả kết quả.',
 'UC19':'Người dùng lựa chọn biểu mẫu hoặc kết quả, kiểm tra dữ liệu và xem trước bố cục trước khi in. Bản dự thảo được nhận diện rõ để tránh sử dụng nhầm với bản đã phát hành.',
 'UC20':'Thông báo được tạo từ sự kiện hoặc lịch nhắc và gửi cho người nhận phù hợp. Hệ thống theo dõi kết quả gửi, sử dụng lại mã thông báo khi thử lại và giới hạn nội dung nhạy cảm trong tin nhắn.',
 'UC21':'Chức năng tra cứu kết hợp từ khóa, bộ lọc và phạm vi quyền của người dùng. Kết quả tìm kiếm và quyền kết xuất được kiểm tra trước khi hiển thị để bảo vệ thông tin của hồ sơ khác.',
 'UC22':'Hướng dẫn giúp người dùng tìm cách thực hiện công việc theo thủ tục và thời điểm áp dụng. Nội dung đi kèm nguồn tham khảo, đồng thời cảnh báo khi tài liệu đã hết hiệu lực để hỗ trợ việc đối chiếu.',
 'UC23':'Báo cáo được lập theo mẫu và kỳ thống kê xác định. Người lập kiểm tra tổng số, các trường hợp ngoại lệ và nguồn hồ sơ trước khi kết xuất để các chỉ tiêu có thể đối chiếu lại.',
 'UC24':'Mỗi mẫu báo cáo quy định công thức tính và các trường hợp loại trừ. Phiên bản mới được kiểm tra bằng dữ liệu mẫu và phê duyệt trước khi công bố, giúp giữ nhất quán số liệu của các kỳ đã chốt.',
 'UC25':'Người dùng chọn chỉ tiêu, cơ quan và thời gian để xem thống kê. Kết quả được liên kết với các hồ sơ tạo nên chỉ tiêu, kèm thông tin về công thức và thời điểm cập nhật để thuận tiện kiểm tra.',
 'UC26':'Màn hình giám sát giúp xác định hồ sơ sắp hạn, quá hạn và cán bộ đang phụ trách. Cảnh báo phục vụ việc đôn đốc, trong khi thời hạn pháp lý tiếp tục được tính theo quy tắc của thủ tục.',
 'UC27':'Chỉ đạo xác định người chịu trách nhiệm, nội dung công việc và hạn phản hồi. Việc theo dõi tiếp nhận, trả lời và kết thúc giúp cơ quan quản lý kiểm tra tiến độ thực hiện theo thẩm quyền.',
 'UC28':'Tài liệu trong kho được phân loại theo nguồn, phiên bản và thời gian hiệu lực. Quyền khai thác được cấp hoặc thu hồi theo mục đích sử dụng, tạo cơ sở bảo vệ và theo dõi việc truy cập tài liệu.',
 'UC29':'Khi dùng lại tài liệu, người nộp lựa chọn nguồn được phép khai thác và kiểm tra nội dung còn phù hợp. Hệ thống lưu tham chiếu để xác định tài liệu đã được dùng cho thành phần nào của hồ sơ.',
 'UC30':'Thông điệp từ hệ thống đối tác được kiểm tra danh tính nguồn, nội dung và thứ tự sự kiện. Dữ liệu hợp lệ được áp dụng theo hợp đồng; trường hợp trùng mã nhưng khác nội dung được giữ lại để đối chiếu.',
 'UC31':'Việc trao đổi với cổng quốc gia sử dụng mã tương quan để theo dõi cùng một yêu cầu qua các lần gửi. Hệ thống ghi nhận biên nhận, kiểm tra kết quả áp dụng và lập lịch gửi lại khi mất phản hồi.',
 'UC32':'Người nộp xem thông báo thu và lựa chọn phương thức thanh toán được hỗ trợ. Giao dịch được liên kết với khoản thu và chứng từ để theo dõi kết quả; thủ tục được miễn hoặc không thu phí sử dụng nhánh xử lý tương ứng.',
 'UC33':'Cán bộ tài chính đối chiếu dữ liệu theo kỳ bằng mã giao dịch và khoản thu. Các chênh lệch được xử lý có căn cứ; đề nghị hoàn trả được phê duyệt và lưu riêng để theo dõi số tiền đã thực hiện.',
 'UC34':'Tài liệu giấy được quét đủ trang, kiểm tra chất lượng và đối chiếu với bản được xuất trình. Bản số hóa được lưu theo phiên bản cùng thông tin xác nhận để sử dụng trong quá trình giải quyết hồ sơ.',
 'UC35':'Người dân gửi nội dung phản ánh hoặc đánh giá và nhận mã theo dõi. Cơ quan được giao tiếp nhận, trả lời hoặc yêu cầu bổ sung; kết quả xử lý phản ánh được quản lý theo quy trình tương ứng.',
 'UC36':'Hồ sơ kết thúc được tập hợp thành gói nộp lưu có danh mục và phiên bản rõ ràng. Hệ thống kiểm tra thành phần, tính toàn vẹn và lưu biên nhận của kho nhận để xác định kết quả bàn giao.',
}

def replace_text(p, text):
    properties=next((deepcopy(r._r.rPr) for r in p.runs if r.text and r._r.rPr is not None),None)
    p.clear()
    run=p.add_run(text)
    if properties is not None:run._r.insert(0,properties)

def polish_report(doc):
    """Sửa phần thân; bỏ qua kết quả trường mục lục do Word tạo."""
    from docx.oxml.ns import qn
    from docx.shared import Pt
    for style in doc.styles:
        if style.name.lower()=='table of figures':
            style.font.name='Times New Roman';style.font.size=Pt(12)
            fmt=style.paragraph_format;fmt.space_before=Pt(0);fmt.space_after=Pt(0);fmt.line_spacing=1.15;fmt.keep_with_next=False
    active_uc=None
    remove_next=False
    changes=0
    for p in list(doc.paragraphs):
        style=p.style.name if p.style else ''
        if style.lower().startswith(('toc','table of')):
            continue
        text=p.text.strip()
        if not text:
            continue
        if style.startswith('Heading'):
            match=re.match(r'^(UC\d{2})\.',text)
            active_uc=match.group(1) if match else None
            remove_next=False
            if text in HEADINGS:
                replace_text(p,HEADINGS[text]);changes+=1
            continue
        if text.startswith('Hướng dẫn thực hiện theo thiết kế:') and active_uc:
            replace_text(p,UC_NARRATIVES[active_uc]);remove_next=True;changes+=1
            continue
        if remove_next:
            # Bảng đặc tả ngay trước đoạn này đã ghi đầy đủ ngoại lệ.
            p._p.getparent().remove(p._p);remove_next=False;changes+=1
            continue
        if text.startswith(('Hình ','Bảng ')):
            continue
        for prefix,new in REVISIONS:
            if text.startswith(prefix):
                if text!=new:replace_text(p,new);changes+=1
                break
    # Khổ A4 dọc thống nhất, giữ lề và kiểu bìa đang sử dụng.
    for section in doc.sections:
        size=section._sectPr.find(qn('w:pgSz'))
        size.set(qn('w:w'),'11906');size.set(qn('w:h'),'16838');size.attrib.pop(qn('w:orient'),None)
    table_revisions={
      '36 hướng dẫn theo bản thiết kế':'Hướng dẫn 36 ca sử dụng, kèm ghi chú về chức năng trong ứng dụng',
      'Tài khoản nội bộ do ứng dụng quản lý; không phải tài khoản công vụ Hà Nội':'Quản lý tài khoản theo vai trò, cơ quan và thời hạn được cấp',
      'Mã Use case của bản thiết kế mới':'Mã ca sử dụng',
    }
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text=cell.text.replace('\u200b','')
                if text in table_revisions:
                    replace_text(cell.paragraphs[0],table_revisions[text])
    return changes
