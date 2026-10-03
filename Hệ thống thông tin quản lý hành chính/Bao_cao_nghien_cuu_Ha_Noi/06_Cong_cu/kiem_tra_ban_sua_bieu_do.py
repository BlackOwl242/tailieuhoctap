from pathlib import Path
import sys,json,re,hashlib
from docx import Document
from PIL import Image,ImageDraw
from lxml import etree as E
ROOT=Path(__file__).resolve().parents[1]
QA=ROOT/'05_Doi_chieu/Sua_bieu_do_va_bien_tap_20261002'
sys.path.insert(0,str(ROOT/'06_Cong_cu'))
import ket_xuat_va_doi_chieu_oo as qa
qa.Q=QA
doc=ROOT/'04_Bao_cao/Phan_tich_thiet_ke_huong_doi_tuong_Ha_Noi.docx'
if sys.argv[1]=='restore':qa.restore(doc)
elif sys.argv[1]=='render':
 qa.render(QA/'Bao_cao.pdf',QA/'Trang')
 pages=sorted((QA/'Trang').glob('trang_*.png'));sheets=QA/'Trang_doi_chieu';sheets.mkdir(exist_ok=True)
 for i in range(0,len(pages),4):
  ims=[Image.open(p) for p in pages[i:i+4]];w,h=ims[0].size
  canvas=Image.new('RGB',(2*w,2*(h+30)),'#dddddd');draw=ImageDraw.Draw(canvas)
  for j,im in enumerate(ims):
   x=(j%2)*w;y=(j//2)*(h+30);draw.text((x+15,y+5),f'TRANG {i+j+1}',fill='black');canvas.paste(im,(x,y+30))
  canvas.save(sheets/f'bo_{i//4+1:02}.png')
 print('Trang:',len(pages),'bộ ảnh:',(len(pages)+3)//4)
elif sys.argv[1]=='check':
 qa.check(doc);d=Document(doc)
 text='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells)
 text=text.replace('\u200b','').replace('\u00ad','')
 checks={
  'Không có mã UCN':not re.search(r'UCN\d+',text),
  'Không có trích dẫn R':not re.search(r'\bR\d+\b',text),
  'Không còn thuật ngữ ca sử dụng':not re.search('ca sử dụng',text,re.I),
  'Không nhắc tới tài liệu mẫu':not any(x in text.lower() for x in ['pttk_oop_hr','talentconnect','đối chiếu mẫu phân tích','tài liệu mẫu','audit mẫu']),
  'Không bàn giao bản tóm tắt':not any(x in text.lower() for x in ['tom_tat_thuyet_trinh','bản tóm tắt']),
  'Đủ 36 biểu đồ trình tự':len([p for p in d.paragraphs if re.search(r'^Hình .*Biểu đồ trình tự của Use case UC\d+',p.text) and '\t' not in p.text])==36,
  'Đủ 36 biểu đồ hoạt động':len([p for p in d.paragraphs if re.search(r'^Hình .*Biểu đồ hoạt động của Use case UC\d+',p.text) and '\t' not in p.text])==36,
 }
 references=set(re.findall(r'\[(\d+)\]',text));checks['Dẫn đầy đủ 10 nguồn']=references==set(str(i) for i in range(1,11))
 figure_checks=[]
 for n in range(1,37):
  f=ROOT/f'03_Thiet_ke/Huong_doi_tuong/So_do/UC{n:02}_trinh_tu.svg';s=f.read_text('utf-8');visible=' '.join(E.fromstring(s.encode('utf-8')).xpath('//*[local-name()="text"]/text()'))
  figure_checks.append({'uc':f'UC{n:02}','read_return':'Dữ liệu và phiên bản' in visible,'domain_return':'Kết quả kiểm tra' in visible,'commit_return':'Xác nhận giao dịch đã lưu' in visible,'actor_result':'Hiển thị kết quả' in visible,'actor_error':'Hiển thị lỗi' in visible,'activation_count':len(re.findall(r'<rect[^>]*width="20"',s))})
 checks['Các luồng có phản hồi, thanh thực thi và kết quả tới tác nhân']=all(all(v for k,v in a.items() if k not in ['uc','activation_count']) and a['activation_count']>=6 for a in figure_checks)
 pages=json.loads((QA/'Trang/van_ban_theo_trang.json').read_text('utf-8'))
 checks['Không có chữ vượt trang']=all(not p['overflow'] for p in pages)
 result={'sha256':hashlib.sha256(doc.read_bytes()).hexdigest(),'pages':len(pages),'checks':checks,'sequences':figure_checks,'summary_deletion':'Bị công cụ từ chối theo chính sách; tệp vẫn còn, không còn liên kết trong báo cáo'}
 (QA/'Kiem_tra_ban_sua.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
 print(json.dumps({'pages':len(pages),'checks':checks},ensure_ascii=False,indent=2))
