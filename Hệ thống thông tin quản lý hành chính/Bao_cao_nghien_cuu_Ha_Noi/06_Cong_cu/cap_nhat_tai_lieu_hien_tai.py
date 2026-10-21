"""Update runtime documentation and the existing research report in place."""
from pathlib import Path
import hashlib, json, re, sys
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))
from noi_dung_van_hanh_hien_tai import DATE, ROOT as PROJECT, COVERAGE, manifest
from chen_noi_dung_hien_tai import add_ui_section, add_coverage_section

def set_text(p, text):
    p.clear()
    p.add_run(text)

def remove_between(doc, start, end):
    body = doc._element.body
    elems = list(body)
    a, b = elems.index(start._p), elems.index(end._p)
    for el in elems[a+1:b]: body.remove(el)

def update_report():
    path = ROOT/'04_Bao_cao/Phan_tich_thiet_ke_huong_doi_tuong_Ha_Noi.docx'
    doc = Document(path)
    heading = lambda p: p.style is not None and p.style.name.startswith('Heading')
    h46 = next(p for p in doc.paragraphs if heading(p) and p.text.strip() == '4.6. Giao diện hệ thống')
    ch5 = next(p for p in doc.paragraphs if heading(p) and p.text.strip() == 'CHƯƠNG 5. HƯỚNG DẪN THEO THỦ TỤC CỤ THỂ')
    remove_between(doc, h46, ch5)
    add_ui_section(doc, anchor=ch5._p)

    # Current implementation notes replace obsolete homepage and guide wording.
    ps = list(doc.paragraphs)
    p620 = next(p for p in ps if p.text.startswith('Ứng dụng React 19, TypeScript, Python và SQLite lưu hồ sơ'))
    set_text(p620, 'Ứng dụng React 19, TypeScript, Python và SQLite lưu hồ sơ, tệp, tài khoản, quyền, khoản thu, lịch hẹn, PAKN và nhật ký trong cơ sở dữ liệu dự án (data/hanoi-services.sqlite3). Giao dịch mới phát sinh từ thao tác người dùng. Luồng nội bộ gồm nộp/tiếp nhận, phân công, thụ lý, trình duyệt, ký nội bộ Ed25519, phát hành, giao nhận, đối chiếu khoản thu và hoàn phí. Lịch hẹn kiểm tra cơ quan, thủ tục, ngày làm việc và sức chứa; phản ánh được phân công theo cơ quan và chỉ công khai sau khi trả lời. Trang chủ đã bỏ khối tra cứu nhanh và số liệu trùng với trang /tra-cuu-ho-so; thống kê vận hành nằm trong bàn làm việc và chỉ phản ánh dữ liệu của ứng dụng.')
    p621 = next(p for p in ps if p.text.startswith(('Cẩm nang (/huong-dan)', 'Cẩm nang /huong-dan')))
    set_text(p621, 'Cẩm nang /huong-dan có hướng dẫn cho công dân/tổ chức và sáu vai trò cán bộ, kèm bước thao tác và nút mở chức năng tương ứng. Lối vào nghiệp vụ cán bộ yêu cầu đăng nhập; tài khoản, quyền và danh mục được quản lý nội bộ, chưa liên kết định danh hoặc danh bạ cán bộ của Thành phố.')
    p622 = next(p for p in ps if p.text.startswith('procedure_forms.json có 2.075 bản ghi'))
    set_text(p622, p622.text + ' Danh mục 36 use case tại mục 6.5 chỉ ra rõ các luồng đã có, một phần và chưa triển khai; không suy diễn đủ số bản ghi biểu mẫu thành đủ cấu hình nghiệp vụ đã được Hà Nội xác nhận.')

    # Target model remains as a design; distinguish it from current application roles.
    p118 = next((p for p in doc.paragraphs if 'quản trị kỹ thuật' in p.text.lower() and ('hồ sơ' in p.text.lower())), None)
    if p118:
        set_text(p118, p118.text + ' Đây là nguyên tắc mục tiêu của mô hình phân tách nhiệm vụ. Trong bản cài đặt hiện tại, vai trò Quản trị hệ thống được cấp quyền xem dữ liệu hồ sơ trong ứng dụng; chưa có vai trò quản trị hạ tầng riêng tách khỏi quyền dữ liệu.')
    for table in doc.tables:
        for row in table.rows:
            vals = [c.text for c in row.cells]
            if any('DA_PHE_DUYET' in x for x in vals):
                for cell in row.cells:
                    if 'Đã ký' in cell.text or 'đã được ký' in cell.text.lower():
                        set_text(cell.paragraphs[0], 'Đã được lãnh đạo phê duyệt; chữ ký là bước riêng sau đó')
            if any('Anh_giao_dien_he_thong' in x.replace('\u200b','') for x in vals):
                for cell in row.cells:
                    if 'cần cập nhật lại ảnh' in cell.text.lower().replace('\u200b',''):
                        set_text(cell.paragraphs[0], f'Đủ 45 ảnh hiện tại, chụp ngày {DATE}; danh mục và SHA-256 ở 03_Thiet_ke/Anh_giao_dien_he_thong/Danh_muc_anh_giao_dien.json')

    # Remove a prior coverage section if this updater is run again.
    ps = list(doc.paragraphs)
    existing = next((p for p in ps if heading(p) and p.text.strip().startswith('6.5. Đối chiếu use case với bản triển khai')), None)
    ch7 = next(p for p in ps if heading(p) and p.text.strip().startswith('CHƯƠNG 7. TRIỂN KHAI'))
    if existing:
        remove_between(doc, existing, ch7)
        existing._p.getparent().remove(existing._p)
    add_coverage_section(doc, anchor=ch7._p)

    # Appendix photo inventory points to the current 45-image capture set.
    for p in doc.paragraphs:
        if 'cần cập nhật lại ảnh' in p.text.lower() or 'ảnh giao diện' in p.text.lower() and '2026' in p.text.lower() and 'Phụ lục' in p.text:
            set_text(p, f'Ảnh giao diện ứng dụng trong mục 4.6 được chụp lại ngày {DATE}, đủ 45 ảnh, có danh mục và tuyến dẫn trong 03_Thiet_ke/Anh_giao_dien_he_thong/Danh_muc_anh_giao_dien.json. Ảnh lịch sử từ nguồn công khai ở chương 2 giữ riêng theo xuất xứ.')
    doc.save(path)

def update_manual():
    out = ROOT/'07_Huong_dan/Huong_dan_van_hanh_he_thong.md'
    roles = [
        ('Công dân hoặc tổ chức','Mở danh mục và chi tiết thủ tục → lập tờ khai năm bước → tải tài liệu thực tế → rà soát → gửi. Lưu mã hồ sơ và khóa tra cứu trên biên nhận. Dùng cả hai tại /tra-cuu-ho-so để xem trạng thái, yêu cầu bổ sung, khoản thu và kết quả.'),
        ('Cán bộ tiếp nhận','Đăng nhập /staff → mở hộp xử lý hoặc tiếp nhận trực tiếp → kiểm tra thành phần → lập Mẫu 01, yêu cầu bổ sung Mẫu 02 hoặc từ chối tiếp nhận Mẫu 03 → định khoản thu nếu có → phân công chuyên viên.'),
        ('Chuyên viên thụ lý','Mở hồ sơ được giao → kiểm tra tài liệu và căn cứ → ghi ý kiến/thẩm định → soạn kết quả → lập Mẫu 04 nếu đủ căn cứ gia hạn → trình lãnh đạo. Khi nhận bổ sung, tiếp tục từ giai đoạn đã lưu.'),
        ('Lãnh đạo phê duyệt','Kiểm tra hồ sơ và dự thảo → phê duyệt hoặc trả lại/từ chối giải quyết. Khi được cấp quyền ký, ký xác nhận nội bộ kết quả đã duyệt. Cùng một lãnh đạo có thể vừa duyệt vừa ký nếu quyền hiện hành cho phép.'),
        ('Văn thư phát hành','Kiểm tra bản đã duyệt, chữ ký nội bộ hợp lệ và khoản thu đã đối chiếu/miễn → cấp số → phát hành → ghi người nhận, kênh giao và chứng cứ. Mã vận đơn nếu có được nhập thủ công.'),
        ('Kế toán hoặc tài chính','Mở sổ khoản thu → đối chiếu mã/chứng từ do người nộp khai → ghi kết quả và tham chiếu hoặc miễn thu. Hoàn trả: lập đề nghị → duyệt theo quyền → ghi chứng từ thực hiện. Đây là sổ nội bộ, không phải chuyển tiền.'),
        ('Quản trị hệ thống','Quản lý tài khoản, vai trò, đơn vị, thời hạn, quyền ký, danh mục, biểu mẫu, tin tức và báo cáo. Tùy quyền, kích hoạt trực tiếp hoặc gửi đề nghị chờ duyệt; kiểm tra lịch sử thay đổi. Vai trò này hiện có thể xem hồ sơ trong ứng dụng.'),
    ]
    text = ['# Hướng dẫn vận hành hệ thống', '', f'Cập nhật {DATE}. Hướng dẫn này mô tả các thao tác đang có trong ứng dụng của dự án và chỉ rõ những kết nối ngoài chưa hoạt động.', '', '## Cách vào hệ thống', '', '- Công dân: `/`, `/thu-tuc`, `/nop-ho-so`, `/tra-cuu-ho-so` và `/huong-dan`.', '- Cán bộ: mở `/staff`, đăng nhập tài khoản nội bộ được cấp, rồi chọn phân hệ theo vai trò. Hướng dẫn ở `/huong-dan` mở đăng nhập nếu chưa có phiên; sau đăng nhập chọn chức năng trên thanh điều hướng.', '- Trang chủ không còn khối tra cứu nhanh và ô số liệu hồ sơ. Tra cứu thực hiện tại `/tra-cuu-ho-so` bằng mã hồ sơ cùng khóa bí mật trên biên nhận.', '', '## Vai trò và thao tác', '']
    for role, steps in roles: text += [f'### {role}', '', steps, '']
    text += ['## Luồng hồ sơ và dữ liệu', '', '1. Gửi hồ sơ tạo trạng thái `CHO_TIEP_NHAN`, mã hồ sơ, khóa tra cứu, phiên bản tờ khai và tệp đính kèm.', '2. Tiếp nhận chuyển sang `DA_TIEP_NHAN`, lưu Mẫu 01 và hạn xử lý; phân công chuyển `DANG_THU_LY`.', '3. Chuyên viên lưu ý kiến/dự thảo và trình `TRINH_DUYET`; lãnh đạo phê duyệt thành `DA_PHE_DUYET` hoặc trả lại.', '4. Khoản thu phải được đối chiếu hoặc miễn trước ký và phát hành. Ký nội bộ kiểm tra quyền, nội dung được duyệt và mã băm; trạng thái vẫn là `DA_PHE_DUYET`.', '5. Văn thư cấp số và phát hành thành `DA_PHAT_HANH`; ghi bằng chứng giao để hoàn tất `HOAN_THANH`.', '', '### Nhánh ngoại lệ', '', '- Bổ sung: Mẫu 02; lưu giai đoạn quay lại và nhận tài liệu mới.', '- Từ chối tiếp nhận: Mẫu 03, trạng thái `TU_CHOI`.', '- Từ chối giải quyết sau thụ lý: `KHONG_GIAI_QUYET`.', '- Gia hạn: Mẫu 04, theo điều kiện máy chủ và căn cứ.', '- Dừng theo yêu cầu: lưu Mẫu 05 đã ký; người có thẩm quyền đính kèm Mẫu 06 đã ký, phát hành rồi chuyển `DUNG`.', '- Lỗi hoặc thiếu kết nối không được ghi như xác nhận từ dịch vụ ngoài.', '', '## Phạm vi chữ ký, thu phí và tích hợp', '', 'Chữ ký Ed25519 bảo vệ dữ liệu kết quả chuẩn hóa trong ứng dụng; không phải chữ ký công vụ trên PDF. Khoản thu/hoàn phí lưu trạng thái và tham chiếu đối chiếu; không kết nối ngân hàng/Kho bạc. Tài khoản cán bộ cấp nội bộ, không đồng bộ danh bạ công vụ. Mã vận đơn, nếu có, được nhập thủ công.', '', 'Cổng Dịch vụ công quốc gia, VNeID, cơ sở dữ liệu dân cư/chuyên ngành, thanh toán/Kho bạc, chứng thư số công vụ, SMS/email, bưu chính, iHanoi và kho lưu trữ ngoài chưa kết nối. Liên kết mở trang ngoài không đồng nghĩa API liên thông.', '', '## Tài liệu liên quan', '', '- [Đối chiếu 36 use case với bản triển khai](../05_Doi_chieu/Doi_chieu_36_use_case_hien_tai.md)', '- [Danh sách tài khoản được phân quyền](../../DANH_SACH_TAI_KHOAN_DUOC_PHAN_QUYEN.md)', '- Hướng dẫn theo vai trò trên giao diện: `/huong-dan`.', '']
    out.write_text('\n'.join(text), encoding='utf-8')

def update_coverage_and_uc():
    out = ROOT/'05_Doi_chieu/Doi_chieu_36_use_case_hien_tai.md'
    lines = [f'# Đối chiếu 36 use case với bản triển khai hiện tại', '', f'Cập nhật {DATE}. Các UC mô tả thiết kế mục tiêu; bảng dưới ghi nhận hiện trạng ứng dụng React/TypeScript và Python/SQLite. “Đã có luồng nội bộ” không khẳng định tích hợp, nghiệm thu hoặc vận hành trên hạ tầng Thành phố.', '', '| use case | Mức hiện thực | Màn hình | Phạm vi và giới hạn |', '|---|---|---|---|']
    for r in COVERAGE:
        clean = lambda s: str(s).replace('|','\\|').replace('\n',' ')
        lines.append('| '+' | '.join(map(clean,[r['id']+' '+r['name'],r['status'],r['ui'],r['note']]))+' |')
    out.write_text('\n'.join(lines)+'\n', encoding='utf-8')
    folder = ROOT/'07_Huong_dan/Chuc_nang_dung_chung'
    for r in COVERAGE:
        p = folder/f"{r['id']}.md"
        if not p.exists(): continue
        text = p.read_text(encoding='utf-8')
        marker = '\n## Đối chiếu với bản cài đặt cục bộ'
        if marker in text: text = text.split(marker)[0].rstrip()
        # Append exact runtime assessment, keeping the target UC specification intact.
        status = f"\n\n## Đối chiếu với bản cài đặt cục bộ ({DATE})\n\n**Mức hiện thực:** {r['status']}. **Màn hình/thành phần:** {r['ui']}. {r['note']} Đây là đối chiếu với ứng dụng của dự án; các bước còn lại trong đặc tả UC tiếp tục là yêu cầu thiết kế, không phải chức năng đã triển khai.\n"
        p.write_text(text.rstrip()+status, encoding='utf-8')

def update_readmes_and_sources():
    p=ROOT/'07_Huong_dan/README.md';s=p.read_text(encoding='utf-8')
    if 'Huong_dan_van_hanh_he_thong.md' not in s:
        s += '\n6. [Hướng dẫn vận hành hiện tại](Huong_dan_van_hanh_he_thong.md) mô tả luồng thực thi theo vai trò, trạng thái, dữ liệu và giới hạn tích hợp. Trạng thái từng UC: [đối chiếu 36 use case](../05_Doi_chieu/Doi_chieu_36_use_case_hien_tai.md).\n'
    p.write_text(s,encoding='utf-8')
    p=ROOT/'03_Thiet_ke/README.md';s=p.read_text(encoding='utf-8')
    if 'Danh_muc_anh_giao_dien.json' not in s:s += '\n* `Anh_giao_dien_he_thong/Danh_muc_anh_giao_dien.json`: danh mục 45 ảnh giao diện ứng dụng, chụp ngày '+DATE+'; không phải ảnh hệ thống nội bộ Hà Nội.\n* Hiện trạng ứng dụng và mức bao phủ 36 use case: `../05_Doi_chieu/Doi_chieu_36_use_case_hien_tai.md`.\n'
    p.write_text(s,encoding='utf-8')
    p=ROOT/'README.md';s=p.read_text(encoding='utf-8')
    s=s.replace('Các ảnh giao diện được lưu trong `03_Thiet_ke/Anh_giao_dien_he_thong`. Khi giao diện có thay đổi, cần chụp lại ảnh trước khi đóng gói tài liệu để nội dung và ảnh cùng phiên bản.', f"Báo cáo Word đã được đồng bộ ngày {DATE} với 45 ảnh giao diện hiện tại, mục đối chiếu đủ 36 use case và cẩm nang vận hành. Xem `07_Huong_dan/Huong_dan_van_hanh_he_thong.md`, `05_Doi_chieu/Doi_chieu_36_use_case_hien_tai.md` và danh mục ảnh `03_Thiet_ke/Anh_giao_dien_he_thong/Danh_muc_anh_giao_dien.json`. Ảnh được chụp từ ứng dụng của dự án; khi giao diện đổi, cập nhật danh mục ảnh và tài liệu cùng phiên bản.")
    p.write_text(s,encoding='utf-8')
    p=ROOT/'02_Noi_dung/05_Trien_khai_van_hanh.md';s=p.read_text(encoding='utf-8')
    s=s.replace('Quản trị viên hệ thống là đầu mối giao quyền quản lý tài khoản theo cơ quan qua đề nghị có căn cứ và phê duyệt độc lập. Người được giao chỉ được xem và lập yêu cầu trong đúng cơ quan; không tự phê duyệt, chuyển người sang đơn vị khác hoặc quản lý tài khoản quản trị hệ thống. Cần rà soát quyền theo quyết định phân công, thời hạn công tác và chuyển đơn vị; khi duyệt đổi quyền, kết thúc phiên cũ và lưu nhật ký trước/sau. Dữ liệu hiện có của ứng dụng chưa có tài khoản nào được giao quyền quản lý cơ quan; đây là trạng thái khởi tạo cục bộ, không phản ánh cách tổ chức tài khoản của Hà Nội.', 'Trong bản cài đặt hiện tại, quản trị viên có thể kích hoạt tài khoản trực tiếp hoặc gửi đề nghị chờ duyệt theo quyền được cấp. Việc thay đổi vai trò, đơn vị, hiệu lực và quyền ký được lưu trong nhật ký; tài khoản có thể bị khóa/thu hồi phiên. Đây là quản trị cục bộ của ứng dụng, không xác định mô hình phân quyền chính thức Hà Nội. Ma trận mục tiêu tiếp tục quy định phân tách nhiệm vụ và phạm vi theo cơ quan để đối chiếu khi tích hợp thực tế.')
    p.write_text(s,encoding='utf-8')
    p=ROOT/'02_Noi_dung/04_Thiet_ke.md';s=p.read_text(encoding='utf-8')
    marker='## Đối chiếu với bản triển khai hiện tại'
    if marker not in s:
        s += f"\n\n{marker} ({DATE})\n\nCác use case và sơ đồ ở phần trên giữ vai trò thiết kế mục tiêu. Ứng dụng hiện tại đã hiện thực nhiều luồng hồ sơ nội bộ nhưng còn UC mới có một phần hoặc chưa triển khai; đối chiếu theo từng mã nằm tại [bảng 36 use case](../05_Doi_chieu/Doi_chieu_36_use_case_hien_tai.md). Ảnh giao diện của phiên bản đang chạy được chụp thành 45 ảnh và liệt kê kèm SHA-256 tại `03_Thiet_ke/Anh_giao_dien_he_thong/Danh_muc_anh_giao_dien.json`. Các ảnh này thuộc ứng dụng của dự án, không phải giao diện nội bộ của hệ thống Thành phố.\n"
    p.write_text(s,encoding='utf-8')
    p=ROOT/'02_Noi_dung/09_Phu_luc.md';s=p.read_text(encoding='utf-8')
    s=s.replace('`TU_CHOI_TIEP_NHAN`','`TU_CHOI`').replace('`DA_DUNG_THEO_YEU_CAU`','`DUNG`').replace('ký số thử nghiệm','ký xác nhận nội bộ Ed25519 trên dữ liệu kết quả đã duyệt').replace('Trạng thái `DA_PHE_DUYET`','Trạng thái `DA_PHE_DUYET` sau phê duyệt; chữ ký là bước riêng')
    p.write_text(s,encoding='utf-8')

def update_manifest_hashes():
    path=ROOT/'03_Thiet_ke/Anh_giao_dien_he_thong/Danh_muc_anh_giao_dien.json'
    data=json.loads(path.read_text(encoding='utf-8'))
    for r in data['captures']:
        f=path.parent/r['file'];r['sha256']=hashlib.sha256(f.read_bytes()).hexdigest()
    path.write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

if __name__=='__main__':
    update_manifest_hashes()
    update_manual()
    update_coverage_and_uc()
    update_readmes_and_sources()
    update_report()
    print('Updated report, manual, 36 UC guides, source notes, and 45-image manifest.')
