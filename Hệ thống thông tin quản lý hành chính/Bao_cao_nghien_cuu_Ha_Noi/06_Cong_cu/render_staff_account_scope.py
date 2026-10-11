"""Render only the revised report section for visual review; leave the deliverable unchanged."""
import os
import sys
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
import pypdfium2 as pdfium

ROOT = Path(__file__).resolve().parents[1]
QA = ROOT / '05_Doi_chieu/qa_staff_account_20261011'
doc = Document(ROOT / '04_Bao_cao/Phan_tich_thiet_ke_huong_doi_tuong_Ha_Noi.docx')
start = next(p._p for p in doc.paragraphs if p.text.startswith('4.6. Giao diện'))
stop = next(p._p for p in doc.paragraphs if p.text.startswith('CHƯƠNG 5.'))
body = doc.element.body
nodes = list(body)
a, b = nodes.index(start), nodes.index(stop)
for index, node in enumerate(nodes):
    if node.tag != qn('w:sectPr') and not a <= index < b:
        body.remove(node)
source = QA / 'phan_giao_dien_kiem_tra.docx'
doc.save(source)
sys.path.insert(0, str(Path(os.environ['TEMP']) / 'codex-doc-render-aspose'))
import aspose.words as aw
render = aw.Document(str(source))
pdf_path = QA / 'phan_giao_dien_kiem_tra.pdf'
render.save(str(pdf_path))
pdf = pdfium.PdfDocument(pdf_path)
for index in range(len(pdf)):
    pdf[index].render(scale=1.6).to_pil().save(QA / f'page-{index + 1:02}.png')
last_text = pdf[len(pdf) - 1].get_textpage().get_text_range()
assert 'quản lý phiên đăng nhập' in last_text, last_text
print(f'Rendered all {len(pdf)} pages of the revised UI section. Final report was not processed by the evaluation renderer.')
