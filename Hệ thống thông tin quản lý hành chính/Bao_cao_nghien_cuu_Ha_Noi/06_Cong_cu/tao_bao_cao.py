from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
from collections import Counter
from hashlib import sha256
import re, json, sys
from docx import Document
from docx.shared import Twips, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path(__file__).resolve().parents[1]
REF=ROOT.parents[1]/'Công nghệ phần mềm/Bao_Cao_Bai_Tap_Lon_CNPM_TalentConnect.docx'
OUT=ROOT/'04_Bao_cao';OUT.mkdir(parents=True,exist_ok=True)
FINAL=OUT/'Bao_cao_nghien_cuu_he_thong_hanh_chinh_Ha_Noi.docx'
assert sha256(REF.read_bytes()).hexdigest()=='256c1192480f9dd050c6c6bd31aea19f959d1c63fa104be20f1e0a332a118725'
d=Document(REF)
source_paras=[deepcopy(p._p) for p in d.paragraphs]
tblpr=deepcopy(d.tables[1]._tbl.tblPr)
sect=deepcopy(d.sections[-1]._sectPr)
body=d.element.body
for child in list(body):body.remove(child)
for p in source_paras[:51]:body.append(p)
body.append(sect)
def replace_p(p,t):
 for child in list(p):
  if child.tag!=qn('w:pPr'):p.remove(child)
 r=OxmlElement('w:r');tx=OxmlElement('w:t');tx.text=t;r.append(tx);p.append(r)
repl={10:'Nghiên cứu hệ thống thông tin giải quyết',11:'thủ tục hành chính thành phố Hà Nội',35:'Nghiên cứu hệ thống thông tin giải quyết',36:'thủ tục hành chính thành phố Hà Nội',14:'Học phần : Hệ thống thông tin quản lý hành chính',39:'Học phần : Hệ thống thông tin quản lý hành chính',12:'',37:'',15:'GVHD      : Hoàng Minh Ngọc',40:'GVHD      : Hoàng Minh Ngọc',41:'Sinh viên : Lê Quóc Huy   2305HTTB011',42:'',43:'',44:'',45:'',46:''}
for i,t in repl.items():
 p=body[i];rpr=deepcopy(p.find(qn('w:r')).find(qn('w:rPr'))) if p.find(qn('w:r')) is not None and p.find(qn('w:r')).find(qn('w:rPr')) is not None else None
 replace_p(p,t)
 if rpr is not None:p.find(qn('w:r')).insert(0,rpr)
for p in body[:51]:
 for tx in p.iter(qn('w:t')):
  if tx.text:tx.text=tx.text.replace(' – ','   ')
def add_field(p,code,placeholder=''):
 r=OxmlElement('w:r');ch=OxmlElement('w:fldChar');ch.set(qn('w:fldCharType'),'begin');r.append(ch);p._p.append(r)
 r=OxmlElement('w:r');t=OxmlElement('w:instrText');t.set(qn('xml:space'),'preserve');t.text=' '+code+' ';r.append(t);p._p.append(r)
 r=OxmlElement('w:r');ch=OxmlElement('w:fldChar');ch.set(qn('w:fldCharType'),'separate');r.append(ch);p._p.append(r)
 p.add_run(placeholder)
 r=OxmlElement('w:r');ch=OxmlElement('w:fldChar');ch.set(qn('w:fldCharType'),'end');r.append(ch);p._p.append(r)
def heading(t,level=1):
 p=d.add_paragraph(t,style=f'Heading {level}')
 if level==1:p.paragraph_format.page_break_before=True
 if level==3:
  for r in p.runs:r.bold=False;r.italic=True
 return p
def para(t):
 p=d.add_paragraph();last=0
 for m in re.finditer(r'https?://[^\s]+',t):
  p.add_run(t[last:m.start()])
  link=OxmlElement('w:hyperlink');link.set(qn('r:id'),d.part.relate_to(m.group(),RT.HYPERLINK,is_external=True))
  r=OxmlElement('w:r');tx=OxmlElement('w:t');tx.text=m.group();r.append(tx);link.append(r);p._p.append(link);last=m.end()
 p.add_run(t[last:]);return p
def cloned_caption(source,text):
 p=d.add_paragraph();p._p.clear()
 if source_paras[source].find(qn('w:pPr')) is not None:p._p.append(deepcopy(source_paras[source].find(qn('w:pPr'))))
 r=p.add_run(text)
 sr=source_paras[source].find(qn('w:r'))
 if sr is not None and sr.find(qn('w:rPr')) is not None:r._r.insert(0,deepcopy(sr.find(qn('w:rPr'))))
 return p
table_counts=Counter();figure_counts=Counter();caps=[];figcaps=[]
def caption(kind,chapter,text):
 counts=table_counts if kind=='Bảng' else figure_counts
 counts[chapter]+=1;n=counts[chapter]
 p=cloned_caption(244 if kind=='Bảng' else 173,'')
 p.add_run(f'{kind} {chapter}.')
 add_field(p,f'SEQ {kind}'+(' \\r 1' if n==1 else ''),str(n))
 p.add_run(': '+text)
 for r in p.runs:
  if kind=='Bảng':r.bold=True
  else:r.italic=True;r.font.size=Pt(11)
 (caps if kind=='Bảng' else figcaps).append(f'{kind} {chapter}.{n}: {text}')
 return p
def table(rows,widths=None):
 n=len(rows[0]);t=d.add_table(rows=len(rows),cols=n)
 old=t._tbl.tblPr;t._tbl.replace(old,deepcopy(tblpr))
 if widths is None:
  widths={2:[2200,6872],3:[1650,3200,4222],4:[1000,2600,2500,2972],5:[1800,1700,1700,1900,1972]}[n]
 for col,w in zip(t.columns,widths):col.width=Twips(w)
 for ri,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if ri==0:trpr.append(OxmlElement('w:tblHeader'))
  for ci,(cell,text) in enumerate(zip(row.cells,rows[ri])):
   cell.width=Twips(widths[ci]);cell.text=re.sub(r'\b[A-Za-z0-9]+(?:_[A-Za-z0-9]+)+\b',lambda m:m.group().replace('_','_\u200b'),text)
   for p in cell.paragraphs:
    pf=p.paragraph_format;pf.first_line_indent=Twips(0);pf.space_before=Pt(2);pf.space_after=Pt(2);pf.line_spacing=1
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if ri==0 or (ci==0 and len(text)<12) else WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs:r.bold=ri==0
 return t
heading('LỜI CẢM ƠN')
para('Em trân trọng cảm ơn giảng viên Hoàng Minh Ngọc đã hướng dẫn học phần Hệ thống thông tin quản lý hành chính. Những kiến thức của học phần giúp em tiếp cận việc phân tích hệ thống gắn với tổ chức công vụ, xem xét quy trình, dữ liệu, trách nhiệm và điều kiện vận hành.')
para('Em trân trọng ghi nhận các tài liệu được cơ quan nhà nước công khai, đặc biệt là thông tin vận hành của Hà Nội, quy định chức năng của hệ thống giải quyết thủ tục hành chính và các quy trình điện tử chuyên ngành. Các tài liệu này tạo căn cứ để lựa chọn đối tượng thực tế, đối chiếu những thay đổi theo thời gian và xây dựng phương án thiết kế.')
para('Em mong nhận được ý kiến của giảng viên để tiếp tục hoàn thiện lập luận, kiểm chứng các đề xuất và bổ sung khảo sát với đơn vị sử dụng khi có điều kiện.')
heading('LỜI CAM ĐOAN')
para('Em chịu trách nhiệm về việc dẫn nguồn và phân định giữa thông tin công khai, suy luận phân tích và thiết kế đề xuất trong báo cáo. Các số liệu lịch sử được ghi theo thời điểm và phạm vi của tài liệu nguồn; các chỉ tiêu kỹ thuật được nêu là mục tiêu đề xuất khi chưa có kết quả đo kiểm.')
para('Báo cáo không khẳng định đã khảo sát nội bộ, triển khai hoặc kiểm thử trên hệ thống vận hành của thành phố Hà Nội. Sơ đồ, mô hình dữ liệu và giao tiếp được xây dựng để phục vụ phân tích, thẩm định và phát triển tiếp theo. Những nội dung còn cần xác nhận được nêu trong phạm vi nghiên cứu và phụ lục khảo sát.')
for title,field in [('MỤC LỤC','TOC \\o "1-2" \\h \\z'),('DANH MỤC BẢNG BIỂU','TOC \\h \\z \\c "Bảng"'),('DANH MỤC HÌNH VÀ SƠ ĐỒ','TOC \\h \\z \\c "Hình"')]:
 p=cloned_caption(116,title);p.paragraph_format.page_break_before=True
 p=d.add_paragraph();p.paragraph_format.first_line_indent=Twips(0);add_field(p,field,'Cập nhật trường trong Microsoft Word')
p=cloned_caption(116,'DANH MỤC TỪ VÀ THUẬT NGỮ VIẾT TẮT');p.paragraph_format.page_break_before=True
table([['STT','Từ viết tắt','Diễn giải'],['1','TTHC','Thủ tục hành chính'],['2','DVC','Dịch vụ công'],['3','UBND','Ủy ban nhân dân'],['4','CSDL','Cơ sở dữ liệu'],['5','API','Giao diện lập trình ứng dụng'],['6','PK','Khóa chính'],['7','FK','Khóa ngoại'],['8','UQ','Ràng buộc duy nhất'],['9','RTO','Mục tiêu thời gian phục hồi'],['10','RPO','Mục tiêu điểm phục hồi'],['11','WCAG','Hướng dẫn khả năng tiếp cận nội dung web'],['12','OWASP','Tổ chức cung cấp tài liệu tham chiếu an toàn ứng dụng'],['13','ASVS','Bộ yêu cầu kiểm chứng an toàn ứng dụng'],['14','VNeID','Ứng dụng định danh điện tử quốc gia']],widths=[760,1900,6412])

full=[]
for file in sorted((ROOT/'02_Noi_dung').glob('*.md')):
 txt=file.read_text(encoding='utf-8').replace('[[DATA_DICTIONARY]]',(ROOT/'03_Thiet_ke/Tu_dien_du_lieu.md').read_text(encoding='utf-8'))
 lines=txt.splitlines();i=0;chapter=str(int(file.name[:2])) if file.name[:2]<'08' else 'A'
 while i<len(lines):
  line=lines[i].strip()
  if not line:i+=1;continue
  if line.startswith('# '):heading(line[2:]);full.append(line)
  elif line.startswith('## '):heading(line[3:],2);full.append(line)
  elif line.startswith('### '):heading(line[4:],3);full.append(line)
  elif re.match(r'Bảng [\dA-Z]+\.\d+:',line):
   prefix=re.match(r'Bảng ([\dA-Z]+)\.',line).group(1)
   c=chapter if prefix.isdigit() else prefix
   title=line.split(': ',1)[1];caption('Bảng',c,title)
   full.append(caps[-1])
  elif line.startswith('|'):
   rows=[]
   while i<len(lines) and lines[i].strip().startswith('|'):
    cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
    if not all(re.fullmatch(r':?-+:?',x) for x in cells):rows.append(cells)
    full.append(lines[i]);i+=1
   widths=None
   if rows[0][0]=='Tên trường':widths=[2200,1750,1450,1100,2572]
   if rows[0][0]=='Trạng thái' and len(rows[0])==3:widths=[3000,2100,3972]
   table(rows,widths);continue
  elif re.match(r'Hình [\dA-Z]+\.\d+:',line):
   title=line.split(': ',1)[1]
   j=i+1
   while j<len(lines) and not lines[j].strip():j+=1
   m=re.match(r'!\[.*?\]\((.*?)\)',lines[j])
   if not m:raise ValueError('Missing figure '+line)
   p=d.add_paragraph();p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Twips(0);p.paragraph_format.keep_with_next=True
   p.add_run().add_picture(str((file.parent/m.group(1)).resolve()),width=Twips(8500))
   caption('Hình',chapter,title);full.extend([figcaps[-1],lines[j]]);i=j
  elif line.startswith('!['):raise ValueError('Unpaired image')
  else:para(line);full.append(line)
  i+=1
  full.append('')
d.core_properties.title='Nghiên cứu hệ thống thông tin giải quyết thủ tục hành chính thành phố Hà Nội'
d.core_properties.subject='Phân tích, thiết kế, triển khai, vận hành, kiểm thử và quản trị'
d.core_properties.author='Lê Quóc Huy'
d.core_properties.keywords='Hà Nội, một cửa, hộ tịch, thiết kế hệ thống'
settings=d.settings.element
e=settings.find(qn('w:updateFields'))
if e is None:e=OxmlElement('w:updateFields');settings.append(e)
e.set(qn('w:val'),'true')
d.save(FINAL)
# Preserve template package parts untouched by the content substitution.
editable={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml','word/settings.xml','docProps/core.xml','docProps/app.xml'}
with ZipFile(REF) as zr,ZipFile(FINAL) as zf:
 data={n:zf.read(n) for n in zf.namelist()}
 for n in zr.namelist():
  if n not in editable:data[n]=zr.read(n)
with ZipFile(FINAL,'w',ZIP_DEFLATED) as z:
 for n,b in data.items():z.writestr(n,b)
(OUT/'Bao_cao_tong_hop.md').write_text('\n'.join(full),encoding='utf-8')
(ROOT/'05_Doi_chieu/danh_muc_noi_dung.json').write_text(json.dumps({'tables':caps,'figures':figcaps,'output':str(FINAL)},ensure_ascii=False,indent=2),encoding='utf-8')
print('DOCX',FINAL,'tables',len(d.tables),'figures',len(figcaps),'paragraphs',len(d.paragraphs))
