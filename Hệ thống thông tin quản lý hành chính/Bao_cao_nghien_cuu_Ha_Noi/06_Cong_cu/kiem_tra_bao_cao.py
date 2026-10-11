from pathlib import Path
from zipfile import ZipFile
from hashlib import sha256
from collections import Counter
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
from PIL import Image,ImageDraw
import pypdfium2 as pdfium
import json,re,sys
sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path(__file__).resolve().parents[1]
REF=ROOT.parents[1]/'Công nghệ phần mềm/Bao_Cao_Bai_Tap_Lon_CNPM_TalentConnect.docx'
FINAL=ROOT/'04_Bao_cao/Phan_tich_thiet_ke_huong_doi_tuong_Ha_Noi.docx'
folder=Path(sys.argv[1]) if len(sys.argv)>1 else ROOT/'05_Doi_chieu/bao_cao_render_v1'
folder.mkdir(parents=True,exist_ok=True)
d=Document(FINAL);ref=Document(REF)
checks=[]
def check(name,ok,detail=''):checks.append({'criterion':name,'passed':bool(ok),'detail':detail})
def effective(run, attr):
 value=getattr(run,attr)
 if value is not None:return value
 style=run.style
 while style is not None:
  value=getattr(style.font,attr)
  if value is not None:return value
  style=style.base_style
 style=run._parent.style
 while style is not None:
  value=getattr(style.font,attr)
  if value is not None:return value
  style=style.base_style
 return False
check('Giữ nguyên tài liệu tham chiếu',sha256(REF.read_bytes()).hexdigest()=='256c1192480f9dd050c6c6bd31aea19f959d1c63fa104be20f1e0a332a118725')
check('Hai phần tài liệu',len(d.sections)==len(ref.sections)==2)
for i,(a,b) in enumerate(zip(d.sections,ref.sections)):
 for prop in ['page_width','page_height','top_margin','bottom_margin','left_margin','right_margin','header_distance','footer_distance']:
  check(f'Phần {i+1}: {prop}',getattr(a,prop)==getattr(b,prop),str(getattr(a,prop)))
for name in ['Normal','Heading 1','Heading 2','Heading 3']:
 a=d.styles[name];b=ref.styles[name]
 check('Kiểu '+name,a.element.xml==b.element.xml)
check('Không có tiêu đề cấp bốn',not any(p.style.name=='Heading 4' for p in d.paragraphs))
for p in d.paragraphs:
 if p.style.name=='Heading 3':check('Cấp ba không đậm và có chữ nghiêng: '+p.text,all(not effective(r,'bold') and effective(r,'italic') for r in p.runs if r.text))
fields=[x.text or '' for x in d.element.iter(qn('w:instrText'))]
check('Mục lục chỉ hai cấp',any('TOC' in x and '\\o "1-2"' in x for x in fields))
check('Không có mục lục cấp ba',not any(p.style.name=='toc 3' for p in d.paragraphs))
# Word may split a field instruction across runs, especially at accented text.
check('Có danh mục bảng',bool(re.search(r'\\c\s+"?Bảng"?',''.join(fields))))
check('Có danh mục hình',bool(re.search(r'\\c\s+"?Hình"?',''.join(fields))))
text='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
check('Không dấu gạch ngang kép','--' not in text)
check('Không từ đề tài cũ','TalentConnect' not in text and 'tuyển dụng' not in text)
check('Không chú giải tiếng Anh trong ngoặc',not re.search(r'\([^)]*(?:Microservices|Workflow|API Contract|Event.driven|State Machine|Frontend|Backend)[^)]*\)',text,re.I))
check('Không tiêu đề bốn số',not re.search(r'^\d+\.\d+\.\d+\.\d+\.',text,re.M))
check('Không còn trường chưa cập nhật','Cập nhật trường trong Microsoft Word' not in text)
for ti,t in enumerate(d.tables):
 check(f'Bảng {ti+1} giữ đường viền và lề ô',t._tbl.tblPr.find(qn('w:tblBorders')) is not None and t._tbl.tblPr.find(qn('w:tblCellMar')) is not None)
 check(f'Bảng {ti+1} lặp hàng tiêu đề',bool(t.rows[0]._tr.xpath('./w:trPr/w:tblHeader')))
 check(f'Bảng {ti+1} không chia hàng qua trang',all(r._tr.xpath('./w:trPr/w:cantSplit') for r in t.rows))
 check(f'Bảng {ti+1} không vượt lề',sum(c.width or 0 for c in t.rows[0].cells)<=d.sections[-1].page_width-d.sections[-1].left_margin-d.sections[-1].right_margin+1000)
schema=json.loads((ROOT/'03_Thiet_ke/Tu_dien_du_lieu.json').read_text(encoding='utf-8'))
check('Từ điển có đủ 37 thực thể',len(schema)==37)
check('Cột rỗng của từ điển đủ rộng',all(t.rows[0].cells[3].width>=698500 for t in d.tables if t.cell(0,0).text=='Tên trường'))
check('Ma trận quyền có các cột đọc được',all(min(c.width for c in t.rows[0].cells)>=914400 for t in d.tables if t.cell(0,0).text=='Vai trò'))
check('Khóa ngoại có bảng đích',all(key[3:] in schema for _,cols in schema.values() for _,_,key,_,_ in cols if key.startswith('FK ')))
openapi=json.loads((ROOT/'03_Thiet_ke/Giao_tiep_OpenAPI.json').read_text(encoding='utf-8'))
check('Có 10 giao tiếp tham chiếu',len(openapi['paths'])==10)
import csv
with (ROOT/'03_Thiet_ke/Kich_ban_kiem_thu.csv').open(encoding='utf-8-sig',newline='') as source:cases=list(csv.DictReader(source))
check('Có 101 ca và không ghi kết quả thử chưa thực hiện',len(cases)==101 and all(c['Trạng thái']=='Chưa thực hiện' for c in cases))
with (ROOT/'03_Thiet_ke/Huong_doi_tuong/Kich_ban_kiem_thu_huong_doi_tuong.csv').open(encoding='utf-8-sig',newline='') as source:oo_cases=list(csv.DictReader(source))
check('Có 180 ca hướng đối tượng và chưa ghi kết quả thử',len(oo_cases)==180 and all(c['status']=='Chưa thực hiện' for c in oo_cases))
class_map=json.loads((ROOT/'03_Thiet_ke/Huong_doi_tuong/Anh_xa_lop.json').read_text(encoding='utf-8'))
associations=json.loads((ROOT/'03_Thiet_ke/Huong_doi_tuong/Quan_he_lop_day_du.json').read_text(encoding='utf-8'))
check('Mô hình có 37 lớp và 74 quan hệ',len(class_map['classes'])==37 and len(associations)==74)
appendix_text='\n'.join(cell.text for table in d.tables for row in table.rows for cell in row.cells)
check('Phụ lục phân biệt hai bộ kiểm thử','bộ 180 ca kiểm thử hướng đối tượng' in appendix_text and 'Bộ kiểm thử tổng hợp 101 ca' in appendix_text and '96 ca' not in text and '168 ca kiểm thử dự kiến' not in appendix_text)
cover='\n'.join(p.text for p in d.paragraphs[:51])
check('Thông tin cá nhân đã được gỡ khỏi bìa','Hoàng Minh Ngọc' not in cover and 'Lê Quốc Huy' not in cover and '2305HTTB011' not in cover)
check('Thuộc tính tác giả không chứa tên cá nhân',not d.core_properties.author)
facets=(ROOT/'02_Noi_dung/07_Quan_tri_ket_luan.md').read_text(encoding='utf-8')
check('Có ma trận hai mươi khía cạnh','Hai mươi khía cạnh của toàn hệ thống' in facets and '| 20. Hiệu quả và quản trị |' in facets)
ext=json.loads((ROOT/'03_Thiet_ke/Giao_tiep_mo_rong_OpenAPI.json').read_text(encoding='utf-8'))
check('Bảy giao tiếp mở rộng',sum(len(v) for v in ext['paths'].values())==7)
references=(ROOT/'02_Noi_dung/08_Tai_lieu_tham_khao.md').read_text(encoding='utf-8')
declared=set(re.findall(r'^\[(\d+)\]',references,re.M))
used=set(re.findall(r'\[(\d+)\]',text))
check('Danh mục nghiên cứu khớp đủ 45 nguồn',declared=={str(i) for i in range(1,46)})
word_references=set(re.findall(r'^\[(\d+)\]','\n'.join(p.text for p in d.paragraphs),re.M))
check('Viện dẫn trong báo cáo Word có mục tham khảo tương ứng',len(word_references)>=15 and used<=word_references)
check('Chủ thể hộ tịch chỉ bắt buộc có điều kiện',any(f[0]=='chu_the_ho_tich_id' and f[3]=='Có' for f in schema['ho_so'][1]))
check('Không trùng mã ca',len({c['Mã'] for c in cases})==len(cases))
check('Có ca cho đủ 31 chức năng',{c['Mã'] for c in cases if c['Mã'].startswith('CF')}=={f'CF{i:02}' for i in range(1,32)})
requirements=(ROOT/'02_Noi_dung/03_Phan_tich_yeu_cau.md').read_text(encoding='utf-8')
function_codes=set(re.findall(r'CN\d+',requirements))
check('Đặc tả đủ 31 chức năng',function_codes=={f'CN{i:02}' for i in range(1,32)})
check('Trường dữ liệu không trùng trong thực thể',all(len({f[0] for f in cols})==len(cols) for _,cols in schema.values()))
check('Có cấu hình quy trình theo phiên bản',any(f[0]=='cau_hinh_quy_trinh' for f in schema['phien_ban_thu_tuc'][1]))
links=[rel.target_ref for rel in d.part.rels.values() if rel.reltype.endswith('/hyperlink')]
check('Đường dẫn hệ thống công khai dùng được trong tài liệu',{'https://dichvucong.gov.vn/','https://ttpvhcc.hanoi.gov.vn/trang-chu','https://dangky.dichvucong.gov.vn/'}<=set(links))
check('Tên đơn vị biên soạn đã đối chiếu nguồn','Công ty Cổ phần FPT' not in text and 'Công ty TNHH Hệ thống thông tin FPT' in text)
check('Báo cáo có đối chiếu cổng công khai và không ghi mốc truy cập','VNeID' in text and '2.000908' in text and 'Mốc quan sát cổng công khai' not in text)
pdf=folder/'bao_cao.pdf'
if pdf.exists():
 doc=pdfium.PdfDocument(str(pdf));reader=PdfReader(pdf)
 texts=[p.extract_text() or '' for p in reader.pages]
 check('PDF không có trang trắng',all(x.strip() for x in texts))
 check('Bìa chính đúng đề tài','Hệ thống giải quyết thủ tục hành chính Hà Nội' in texts[0])
 print('PAGES',len(doc))
 (folder/'van_ban_theo_trang.json').write_text(json.dumps(texts,ensure_ascii=False,indent=2),encoding='utf-8')
 for i,page in enumerate(doc):page.render(scale=1.6).to_pil().save(folder/f'page-{i+1:03}.png')
 for start in range(0,len(doc),8):
  im=Image.new('RGB',(1400,1040),'#d6d6d6');dr=ImageDraw.Draw(im)
  for j in range(start,min(start+8,len(doc))):
   # Four columns, two rows, for overview only; page images are retained for reading.
   x=((j-start)%4)*350+10;y=((j-start)//4)*510+25
   page=Image.open(folder/f'page-{j+1:03}.png');page.thumbnail((330,475));im.paste(page,(x,y));dr.text((x,y-18),str(j+1),fill='black')
  im.save(folder/f'contact-{start+1:03}.png')
 report={'pages':len(doc),'tables':len(d.tables),'figures':sum(bool(re.match(r'^Hình [0-9A-Z]+\.\d+:',p.text)) and '\t' not in p.text for p in d.paragraphs),'checks':checks}
else:report={'checks':checks}
(ROOT/'05_Doi_chieu/ket_qua_kiem_tra.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
failed=[x for x in checks if not x['passed']]
print('CHECKS',len(checks),'FAILED',len(failed))
for x in failed:print(x['criterion'],x['detail'])
