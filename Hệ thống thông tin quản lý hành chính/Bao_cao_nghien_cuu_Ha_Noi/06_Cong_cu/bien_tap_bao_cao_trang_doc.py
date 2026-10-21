"""Biên tập văn phong, thống nhất A4 dọc và cập nhật nguồn tạo báo cáo."""
from pathlib import Path
import json,re
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from noi_dung_bao_cao import ROOT, COVERAGE, UI_TITLES, IMPLEMENTATION_NOTES
from chen_noi_dung_hien_tai import add_ui_section, add_coverage_section, add_implementation_section
from van_phong_bao_cao import HEADINGS, REVISIONS, polish_report, replace_text

REPORT=ROOT/'04_Bao_cao/Phan_tich_thiet_ke_huong_doi_tuong_Ha_Noi.docx'

def heading(doc,prefix):
    return next(p for p in doc.paragraphs if p.style.name.startswith('Heading') and p.text.strip().startswith(prefix))

def clear_between(start,end):
    node=start._p.getnext()
    while node is not None and node is not end._p:
        following=node.getnext();node.getparent().remove(node);node=following

def update_report():
    doc=Document(REPORT)
    h46,ch5=heading(doc,'4.6.'),heading(doc,'CHƯƠNG 5.')
    clear_between(h46,ch5)
    add_ui_section(doc,ch5._p)
    h64,h65,ch7=heading(doc,'6.4.'),heading(doc,'6.5.'),heading(doc,'CHƯƠNG 7.')
    clear_between(h64,h65);h64._p.getparent().remove(h64._p)
    add_implementation_section(doc,h65._p)
    clear_between(h65,ch7);h65._p.getparent().remove(h65._p)
    add_coverage_section(doc,ch7._p)
    count=polish_report(doc)
    # Thuộc tính ở cuối phần thân áp dụng cho toàn bộ nội dung sau hai trang bìa.
    final=doc.sections[-1]._sectPr
    page_number=final.find(qn('w:pgNumType'))
    if page_number is None:page_number=OxmlElement('w:pgNumType');final.append(page_number)
    page_number.set(qn('w:start'),'1')
    for table in doc.tables:
        for row in table.rows:
            cells=row.cells
            label=cells[0].text.replace('\u200b','')
            if 'Anh_giao_dien_he_thong' in label:
                replace_text(cells[1].paragraphs[0],'47 ảnh giao diện ứng dụng, chụp ngày 19 và 20/10/2026; ngày chụp từng ảnh có trong danh mục.')
    doc.save(REPORT)
    print(f'Đã biên tập {count} đoạn; {len(doc.inline_shapes)} hình; {len(doc.sections)} phần, đều khổ A4 dọc.')

def write_coverage():
    lines=['# Kết quả triển khai theo use case','',
           'Bảng sau đối chiếu các chức năng của ứng dụng với 36 use case trong mô hình thiết kế. Cột cuối nêu nội dung đã thực hiện và phần cần tiếp tục bổ sung. Tình trạng kết nối dịch vụ ngoài được ghi riêng theo từng chức năng.','',
           '| use case | Kết quả | Chức năng | Nội dung đáp ứng và phần cần bổ sung |',
           '|---|---|---|---|']
    for r in COVERAGE:
        lines.append('| '+' | '.join([r['id']+' '+r['name'],r['status'],r['ui'],r['note']])+' |')
    (ROOT/'05_Doi_chieu/Doi_chieu_36_use_case_hien_tai.md').write_text('\n'.join(lines)+'\n','utf8')
    for r in COVERAGE:
        path=ROOT/'07_Huong_dan/Chuc_nang_dung_chung'/f"{r['id']}.md"
        text=path.read_text('utf8')
        text=re.split(r'\n## (?:Đối chiếu với bản cài đặt cục bộ|Sử dụng trong ứng dụng|Chức năng trong ứng dụng)',text)[0].rstrip()
        text+=f"\n\n## Sử dụng trong ứng dụng\n\n**{r['status']}** tại mục **{r['ui']}**. {r['note']}\n"
        path.write_text(text,'utf8')

def write_manual():
    text='''# Hướng dẫn vận hành hệ thống

Tài liệu hướng dẫn công dân, tổ chức và cán bộ thực hiện các công việc trong hệ thống, từ chuẩn bị hồ sơ đến trả kết quả. Mỗi hồ sơ có mã riêng, lưu thông tin đã kê khai, tài liệu đính kèm và lịch sử xử lý để các bộ phận tiếp tục công việc.

## Truy cập hệ thống

Người dân mở cổng dịch vụ công, chọn **Thủ tục hành chính**, **Nộp hồ sơ trực tuyến**, **Tra cứu hồ sơ** hoặc **Hướng dẫn sử dụng**. Cán bộ chọn **Đăng nhập cán bộ**, dùng tài khoản được quản trị viên cấp và mở chức năng tương ứng trên thanh điều hướng. Khi chưa đăng nhập, các liên kết nghiệp vụ trong cẩm nang sẽ đưa đến trang đăng nhập.

## Công dân và tổ chức

1. Tìm thủ tục theo tên hoặc mã, chọn cơ quan tiếp nhận và đọc thành phần hồ sơ, thời hạn, khoản thu.
2. Mở biểu mẫu nộp hồ sơ, kê khai thông tin người nộp và nội dung theo thủ tục.
3. Tải tài liệu cho từng thành phần, chọn phương thức nhận kết quả và kiểm tra lại thông tin.
4. Xác nhận gửi và giữ mã hồ sơ cùng khóa tra cứu trên biên nhận.
5. Mở trang **Tra cứu hồ sơ**, nhập cả mã và khóa để xem tiến độ, bổ sung tài liệu, khai chứng từ thu và xem kết quả.

Người dân cũng có thể đặt lịch tại Bộ phận Một cửa, gửi phản ánh kiến nghị và theo dõi phản hồi. Đánh giá hồ sơ được thực hiện theo điều kiện hiển thị trên trang tra cứu.

## Cán bộ tiếp nhận

Cán bộ mở **Hộp xử lý** để kiểm tra hồ sơ trực tuyến hoặc chọn **Tiếp nhận trực tiếp** để ghi hồ sơ tại quầy và tải bản quét. Sau khi đối chiếu chủ thể, thành phần và thẩm quyền, cán bộ tiếp nhận lập Mẫu 01, yêu cầu bổ sung bằng Mẫu 02 hoặc từ chối tiếp nhận bằng Mẫu 03. Hồ sơ đủ điều kiện được xác định hạn xử lý, khoản thu nếu có và chuyên viên phụ trách.

## Chuyên viên thụ lý

Chuyên viên mở hồ sơ được giao, kiểm tra tài liệu, ghi ý kiến thẩm định và chuẩn bị dự thảo kết quả. Khi cần bổ sung, nội dung yêu cầu phải chỉ rõ tài liệu và căn cứ để người nộp thực hiện. Sau khi tài liệu được kiểm tra, hồ sơ tiếp tục từ giai đoạn đã lưu. Việc gia hạn sử dụng Mẫu 04 theo điều kiện của thủ tục. Dự thảo hoàn chỉnh được trình lãnh đạo xem xét.

## Lãnh đạo

Lãnh đạo kiểm tra dự thảo cùng căn cứ để phê duyệt, trả lại chỉnh sửa hoặc từ chối giải quyết. Sau khi phê duyệt, tài khoản được cấp quyền ký có thể ký dữ liệu kết quả trong ứng dụng. Một lãnh đạo có thể thực hiện cả duyệt và ký nếu có đủ quyền còn hiệu lực. Khoản thu phải được đối chiếu hoặc miễn trước khi ký và phát hành.

## Văn thư

Văn thư kiểm tra bản đã duyệt, chữ ký và điều kiện khoản thu, sau đó cấp số và phát hành. Khi trả kết quả, cán bộ ghi người nhận, phương thức giao và chứng cứ để hoàn tất hồ sơ. Với kết quả gửi bưu chính, mã vận đơn được nhập để theo dõi thủ công.

## Kế toán và tài chính

Cán bộ tài chính mở **Sổ khoản thu**, đối chiếu mã và chứng từ do người nộp cung cấp rồi ghi kết quả cùng tham chiếu. Trường hợp miễn thu cần lưu căn cứ. Nếu phát sinh hoàn trả, người được giao lập đề nghị, chuyển phê duyệt và ghi chứng từ thực hiện. Sổ này theo dõi nghĩa vụ và bằng chứng thu, hoàn phí trong ứng dụng; việc xác nhận giao dịch tự động phụ thuộc kết nối với đơn vị tài chính.

## Quản trị viên

Quản trị viên cấp hoặc khóa tài khoản, điều chỉnh vai trò, đơn vị, thời hạn và quyền ký. Tùy quyền được giao, tài khoản được kích hoạt trực tiếp hoặc chuyển đề nghị chờ duyệt. Các thay đổi được lưu trong lịch sử. Quản trị viên còn quản lý danh mục, cấu hình biểu mẫu, tin tức và báo cáo; vai trò này hiện được cấp quyền xem dữ liệu hồ sơ trong ứng dụng.

## Các tình huống phát sinh

- **Bổ sung hồ sơ:** lập Mẫu 02, nhận và kiểm tra tài liệu, sau đó tiếp tục xử lý từ giai đoạn đã lưu.
- **Từ chối tiếp nhận:** lập Mẫu 03 và lưu lý do; hồ sơ chuyển sang `TU_CHOI`.
- **Từ chối giải quyết sau thụ lý:** ghi kết luận và căn cứ; hồ sơ chuyển sang `KHONG_GIAI_QUYET`.
- **Gia hạn:** lập Mẫu 04 theo điều kiện và căn cứ áp dụng, lưu hạn cùng lịch sử điều chỉnh.
- **Dừng theo yêu cầu:** lưu Mẫu 05 đã ký, kiểm tra điều kiện và lập thông báo. Người có thẩm quyền tải Mẫu 06 đã ký, ghi nhận phát hành rồi chuyển hồ sơ sang `DUNG`. Trả giấy tờ và hoàn phí được theo dõi riêng.

## Trạng thái hồ sơ

Luồng thông thường gồm `CHO_TIEP_NHAN`, `DA_TIEP_NHAN`, `DANG_THU_LY`, `TRINH_DUYET`, `DA_PHE_DUYET`, `DA_PHAT_HANH` và `HOAN_THANH`. Bước ký được ghi nhận trên nội dung đã phê duyệt, trước khi văn thư phát hành. Mỗi lần chuyển bước đều được kiểm tra theo vai trò, cơ quan và tình trạng hồ sơ.

## Phạm vi kết nối

Ứng dụng sử dụng chữ ký Ed25519 để kiểm tra tính toàn vẹn của dữ liệu kết quả đã duyệt. Các kết nối với Cổng Dịch vụ công quốc gia, VNeID, cơ sở dữ liệu dân cư và chuyên ngành, ngân hàng hoặc Kho bạc, chứng thư số công vụ, SMS, thư điện tử, bưu chính, iHanoi và kho lưu trữ ngoài chưa được triển khai. Việc vận hành chính thức cần có thông tin kết nối và kiểm tra với từng đơn vị quản lý.

## Tài liệu liên quan

- [Kết quả triển khai theo 36 use case](../05_Doi_chieu/Doi_chieu_36_use_case_hien_tai.md)
- [Quản lý tài khoản và phân quyền](../../DANH_SACH_TAI_KHOAN_DUOC_PHAN_QUYEN.md)
- Cẩm nang trên giao diện: `/huong-dan`; bàn làm việc cán bộ: `/staff`.
'''
    (ROOT/'07_Huong_dan/Huong_dan_van_hanh_he_thong.md').write_text(text,'utf8')

def update_sources():
    folder=ROOT/'06_Cong_cu'
    generator=folder/'tao_tai_lieu_huong_doi_tuong.py'
    s=generator.read_text('utf8')
    if 'from van_phong_bao_cao import' not in s:
        s='from van_phong_bao_cao import polish_report, UC_NARRATIVES\n'+s
    s=s.replace("b.p('Hướng dẫn thực hiện theo thiết kế: ' + ' '.join(f'Bước {i}: {s}.' for i,s in enumerate(u['steps'],1)))\n b.p(u['exception']+'.')", "b.p(UC_NARRATIVES[u['id']])")
    s=s.replace('from chen_noi_dung_hien_tai import add_ui_section, add_coverage_section','from chen_noi_dung_hien_tai import add_ui_section, add_coverage_section, add_implementation_section')
    a=s.find("b.h('6.4.")
    b=s.find('add_coverage_section(b.d)',a)
    if a>=0 and b>=0:s=s[:a]+'add_implementation_section(b.d)\n\n'+s[b:]
    if 'polish_report(b.d)' not in s:s=s.replace('b.finish()','polish_report(b.d)\nb.finish()')
    for old,new in HEADINGS.items():s=s.replace(old,new)
    generator.write_text(s,'utf8')
    # Các công cụ cập nhật dùng cùng một bộ nội dung và cùng quy tắc trình bày.
    updater=folder/'cap_nhat_tai_lieu_hien_tai.py';s=updater.read_text('utf8')
    for function,next_function,body in [
      ('update_report','update_manual','    from bien_tap_bao_cao_trang_doc import update_report as edit_report\n    edit_report()\n\n'),
      ('update_manual','update_coverage_and_uc','    from bien_tap_bao_cao_trang_doc import write_manual\n    write_manual()\n\n'),
      ('update_coverage_and_uc','update_readmes_and_sources','    from bien_tap_bao_cao_trang_doc import write_coverage\n    write_coverage()\n\n'),
    ]:
        a=s.index('def '+function+'():');b=s.index('def '+next_function+'():',a)
        s=s[:a]+'def '+function+'():\n'+body+s[b:]
    updater.write_text(s,'utf8')
    # Đồng bộ các đoạn thuyết minh trùng trong bản Markdown.
    for path in (ROOT/'02_Noi_dung').glob('*.md'):
        text=path.read_text('utf8');parts=text.split('\n\n')
        for i,part in enumerate(parts):
            for prefix,new in REVISIONS:
                if part.strip().startswith(prefix):parts[i]=new;break
        text='\n\n'.join(parts)
        for old,new in HEADINGS.items():text=text.replace(old,new)
        marker='## Đối chiếu với bản triển khai hiện tại'
        if marker in text:
            text=text.split(marker)[0].rstrip()+'\n\n## Kết quả triển khai\n\n'+IMPLEMENTATION_NOTES[0]+'\n\nXem [bảng đánh giá theo 36 use case](../05_Doi_chieu/Doi_chieu_36_use_case_hien_tai.md) để đối chiếu chức năng đã hỗ trợ và nội dung cần bổ sung.\n'
        path.write_text(text,'utf8')
    # Chú thích ảnh mô tả chức năng, giữ nguyên tệp ảnh và ngày chụp.
    path=ROOT/'03_Thiet_ke/Anh_giao_dien_he_thong/Danh_muc_anh_giao_dien.json'
    data=json.loads(path.read_text('utf8'))
    capture=folder/'chup_giao_dien_hien_tai.cjs';code=capture.read_text('utf8')
    for n,row in enumerate(data['captures'],start=1):
        old=row['title'];new=UI_TITLES[n]
        code=code.replace(old,new);row['title']=new
    path.write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n','utf8');capture.write_text(code,'utf8')

if __name__=='__main__':
    update_sources()
    write_manual()
    write_coverage()
    update_report()
