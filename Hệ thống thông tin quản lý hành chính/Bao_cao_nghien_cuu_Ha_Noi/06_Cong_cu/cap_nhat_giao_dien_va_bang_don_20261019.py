from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from tempfile import NamedTemporaryFile
from datetime import datetime
import json, re, shutil
from PIL import Image
from docx import Document
from docx.shared import RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
REPORT=ROOT/'04_Bao_cao/Phan_tich_thiet_ke_huong_doi_tuong_Ha_Noi.docx'
ASSET_DIR=ROOT/'03_Thiet_ke/Anh_giao_dien_he_thong'
MANIFEST=json.loads((ASSET_DIR/'Danh_muc_anh_giao_dien.json').read_text(encoding='utf-8'))
assert len(MANIFEST['captures'])==45 and not MANIFEST.get('errors'), MANIFEST.get('errors')
assert all((ASSET_DIR/x['file']).exists() for x in MANIFEST['captures'])
CAPTIONS={26+i:row['title'] for i,row in enumerate(MANIFEST['captures'])}

def set_cell_shading(cell, fill='FFFFFF'):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),fill)

def set_table_borders(table):
    tblPr=table._tbl.tblPr; borders=tblPr.find(qn('w:tblBorders'))
    if borders is None: borders=OxmlElement('w:tblBorders'); tblPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag=qn(f'w:{edge}'); element=borders.find(tag)
        if element is None: element=OxmlElement(f'w:{edge}'); borders.append(element)
        for key,value in [('val','single'),('sz','4'),('space','0'),('color','000000')]: element.set(qn(f'w:{key}'),value)

def replace_paragraph_text(p, text):
    fmt=next((r._r.rPr for r in p.runs if r.text and r._r.rPr is not None),None)
    import copy
    saved=copy.deepcopy(fmt) if fmt is not None else None
    p.clear(); r=p.add_run(text)
    if saved is not None: r._r.insert(0,saved)

doc=Document(REPORT)
# Replace all 45 interface screenshots by the live app captures and preserve the shape's existing width.
media={}; replaced=[]
for shape in doc.inline_shapes:
    inline=shape._inline; docPr=inline.docPr
    match=re.fullmatch(r'Picture (\d+)',docPr.get('name') or '')
    if not match: continue
    number=int(match.group(1))
    if not 163<=number<=207: continue
    idx=number-162; row=MANIFEST['captures'][idx-1]; asset=ASSET_DIR/row['file']
    blip=inline.graphic.graphicData.pic.blipFill.blip
    rid=blip.get(qn('r:embed')); part=doc.part.related_parts[rid]
    media[str(part.partname).lstrip('/')]=asset.read_bytes()
    with Image.open(asset) as im: width,height=im.size
    width_emu=shape.width; shape.height=int(width_emu*height/width)
    docPr.set('descr',row['title']); docPr.set('title',row['title'])
    replaced.append((number,idx,row['title'],width,height))
assert len(replaced)==45, f'Expected 45 embedded UI screenshots, found {len(replaced)}'
# Update figure captions in the body and the cached List of Figures without disturbing page-number fields.
updated_body=updated_list=0
for p in doc.paragraphs:
    match=re.match(r'^Hình 4\.(\d+):',p.text.strip())
    if not match: continue
    number=int(match.group(1))
    if number not in CAPTIONS: continue
    new=f'Hình 4.{number}: {CAPTIONS[number]}'
    if (p.style.name or '').lower()=='table of figures':
        for node in p._p.xpath('.//w:t'):
            if node.text and re.match(r'^Hình 4\.',node.text): node.text=new; updated_list+=1; break
    else:
        if p.text!=new: replace_paragraph_text(p,new); updated_body+=1
# Remove fills and colored text from all tables; use a plain grid with black text and bold first rows.
for table in doc.tables:
    try: table.style='Table Grid'
    except KeyError: pass
    set_table_borders(table)
    for row_idx,row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell,'FFFFFF')
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.color.rgb=RGBColor(0,0,0)
                    rPr=run._r.get_or_add_rPr(); color=rPr.find(qn('w:color'))
                    if color is None: color=OxmlElement('w:color'); rPr.append(color)
                    color.set(qn('w:val'),'000000')
                    for attr in ('themeColor','themeTint','themeShade'): color.attrib.pop(qn('w:'+attr),None)
                    if row_idx==0: run.bold=True
# Correct the screenshot source row wherever it appears in tables.
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            if '45 ảnh giao diện' in cell.text:
                replace_paragraph_text(cell.paragraphs[0],'45 ảnh giao diện ứng dụng, chụp ngày 19/10/2026; có danh mục ảnh kèm theo.')
# Save modified XML, then replace image package parts with the newly captured JPEG bytes.
with NamedTemporaryFile(suffix='.docx',dir=REPORT.parent,delete=False) as tmp: tmp_path=Path(tmp.name)
doc.save(tmp_path)
final_tmp=tmp_path.with_suffix('.packed.docx')
with ZipFile(tmp_path,'r') as src, ZipFile(final_tmp,'w',ZIP_DEFLATED) as dst:
    for item in src.infolist(): dst.writestr(item,media.get(item.filename,src.read(item.filename)))
shutil.move(final_tmp,REPORT); tmp_path.unlink(missing_ok=True)
print(json.dumps({'embeddedScreenshots':len(replaced),'bodyCaptions':updated_body,'figureListEntries':updated_list,'tables':len(doc.tables),'imageParts':len(media),'fileBytes':REPORT.stat().st_size},ensure_ascii=False))
